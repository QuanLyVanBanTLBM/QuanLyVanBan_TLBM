import os
from datetime import date, datetime, time, timedelta
from pathlib import Path

import openpyxl
from django.contrib.auth import get_user_model
from django.contrib.auth.hashers import make_password
from django.contrib.auth.models import Group
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone

from accounts.models import EmployeeAssignment, Position
from authorization.models import GroupInfo
from birthdays.models import BirthdayWish
from meetings.models import MeetingBooking, MeetingRoom
from organizations.models import DonVi


REQUIRED_COLUMNS = {
    "DonVi": {"STT ưu tiên", "Khối", "Tên đơn vị", "Địa chỉ", "Điện thoại", "Công khai"},
    "CBNV": {
        "Tên đăng nhập", "Họ và tên", "Ngày sinh", "Email", "Số điện thoại", "Đơn vị",
        "Khối", "Chức vụ", "Chức danh", "Nhóm quyền", "Trạng thái", "Công khai", "Ghi chú",
    },
    "CBNV_BoSung": {"Email (khóa nối với CBNV_Import)", "Tên đăng nhập", "Nhóm quyền"},
    "KiemNhiem": {
        "Tên đăng nhập", "Khối kiêm nhiệm", "Đơn vị kiêm nhiệm", "Chức vụ kiêm nhiệm",
        "Chức danh kiêm nhiệm",
    },
    "TaiKhoanDemo": {"Tên đăng nhập"},
    "PhongHop": {"Tên phòng", "Vị trí", "Sức chứa", "Tiện ích", "Công khai", "Trạng thái"},
    "DatPhongHop": {
        "Phòng", "Người đặt", "Đơn vị", "Tiêu đề", "Ngày lệch", "Bắt đầu", "Kết thúc",
        "Số người", "Mục đích", "Trạng thái",
    },
    "LoiChucSinhNhat": {
        "Người nhận", "Người gửi", "Nội dung", "Mẫu thiệp", "Ẩn danh", "Trạng thái",
    },
}

EMPLOYEE_STATUS = {"Đang làm việc": "LAM_VIEC", "Nghỉ việc": "NGHI_VIEC"}
ROOM_STATUS = {"Sẵn sàng": "AVAILABLE", "Bảo trì": "MAINTENANCE", "Không khả dụng": "UNAVAILABLE"}
BOOKING_STATUS = {"Đã xác nhận": "CONFIRMED", "Đã hoàn thành": "COMPLETED", "Đã hủy": "CANCELLED"}
WISH_STATUS = {"Đã duyệt": "APPROVED", "Chờ duyệt": "PENDING", "Từ chối": "REJECTED"}


def _text(value):
    return str(value).strip() if value is not None else ""


def _required(record, column, sheet):
    value = _text(record.get(column))
    if not value:
        raise CommandError(f"Sheet {sheet}, dòng {record['__row__']}: thiếu '{column}'.")
    return value


def _parse_bool(value, sheet, row_number, column):
    normalized = _text(value).casefold()
    if normalized == "có":
        return True
    if normalized == "không":
        return False
    raise CommandError(f"Sheet {sheet}, dòng {row_number}: '{column}' phải là Có hoặc Không.")


def _parse_date(value, sheet, row_number, column):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = _text(value)
    try:
        return date.fromisoformat(raw)
    except ValueError as exc:
        raise CommandError(f"Sheet {sheet}, dòng {row_number}: ngày '{column}' không hợp lệ: {raw}.") from exc


def _parse_time(value, sheet, row_number, column):
    if isinstance(value, datetime):
        return value.time()
    if isinstance(value, time):
        return value
    try:
        return time.fromisoformat(_text(value))
    except ValueError as exc:
        raise CommandError(f"Sheet {sheet}, dòng {row_number}: giờ '{column}' không hợp lệ.") from exc


def _parse_int(value, sheet, row_number, column):
    try:
        return int(value)
    except (TypeError, ValueError) as exc:
        raise CommandError(f"Sheet {sheet}, dòng {row_number}: '{column}' phải là số nguyên.") from exc


class Command(BaseCommand):
    help = "Kiểm tra hoặc chủ động nạp workbook dữ liệu mẫu vào DB"

    def add_arguments(self, parser):
        parser.add_argument("--file", required=True, help="Đường dẫn workbook .xlsx")
        parser.add_argument(
            "--apply",
            action="store_true",
            help="Ghi dữ liệu vào DB; mặc định chỉ kiểm tra workbook và đếm dòng.",
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Chỉ kiểm tra workbook, không ghi DB (đây cũng là mặc định).",
        )

    def handle(self, *args, **options):
        if options["apply"] and options["dry_run"]:
            raise CommandError("Chỉ chọn một trong --apply hoặc --dry-run.")

        file_path = Path(options["file"]).expanduser()
        if not file_path.is_file():
            raise CommandError(f"Không tìm thấy file: {file_path}")

        records = self._read_workbook(file_path)
        self._print_counts(records)
        if not options["apply"]:
            self.stdout.write("Kiểm tra xong; chưa ghi dữ liệu vào DB. Dùng --apply để nạp.")
            return

        demo_usernames = {
            _required(row, "Tên đăng nhập", "TaiKhoanDemo")
            for row in records["TaiKhoanDemo"]
        }
        demo_password = os.environ.get("DEMO_PASSWORD", "")
        if demo_usernames and not demo_password.strip():
            raise CommandError("Cần đặt DEMO_PASSWORD để tạo hoặc khởi tạo các tài khoản trong sheet TaiKhoanDemo.")

        with transaction.atomic():
            counts = self._import_records(records, demo_usernames, demo_password)

        self.stdout.write(self.style.SUCCESS("Nạp dữ liệu mẫu hoàn tất."))
        for label, count in counts.items():
            self.stdout.write(f"{label}: {count}")

    def _read_workbook(self, file_path):
        try:
            workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        except Exception as exc:
            raise CommandError(f"Không đọc được file Excel: {exc}") from exc

        try:
            missing_sheets = set(REQUIRED_COLUMNS) - set(workbook.sheetnames)
            if missing_sheets:
                raise CommandError(f"Thiếu sheet: {', '.join(sorted(missing_sheets))}")

            result = {}
            for sheet_name, required_columns in REQUIRED_COLUMNS.items():
                rows = workbook[sheet_name].iter_rows(values_only=True)
                header_row = next(rows, None)
                if header_row is None:
                    raise CommandError(f"Sheet {sheet_name} đang trống.")

                headers = [_text(value) for value in header_row]
                missing_columns = required_columns - set(headers)
                if missing_columns:
                    raise CommandError(
                        f"Sheet {sheet_name} thiếu cột: {', '.join(sorted(missing_columns))}"
                    )

                sheet_records = []
                for row_number, values in enumerate(rows, start=2):
                    if not any(value is not None and _text(value) for value in values):
                        continue
                    record = {header: value for header, value in zip(headers, values) if header}
                    record["__row__"] = row_number
                    sheet_records.append(record)
                result[sheet_name] = sheet_records
            self._validate_workbook_rows(result)
            return result
        finally:
            workbook.close()

    def _validate_workbook_rows(self, records):
        employees_by_username = {}
        employees_by_email = {}
        for row in records["CBNV"]:
            username = _required(row, "Tên đăng nhập", "CBNV")
            email = _required(row, "Email", "CBNV").casefold()
            if username in employees_by_username:
                raise CommandError(f"Sheet CBNV có username trùng: {username}.")
            if email in employees_by_email:
                raise CommandError(f"Sheet CBNV có email trùng: {email}.")
            employees_by_username[username] = row
            employees_by_email[email] = row

            for column in ("Họ và tên", "Ngày sinh", "Khối", "Đơn vị", "Chức vụ", "Trạng thái", "Công khai"):
                _required(row, column, "CBNV")
            if _text(row["Trạng thái"]) not in EMPLOYEE_STATUS:
                raise CommandError(f"Sheet CBNV, dòng {row['__row__']}: trạng thái không hợp lệ.")
            _parse_bool(row["Công khai"], "CBNV", row["__row__"], "Công khai")

        role_by_email = {}
        for row in records["CBNV_BoSung"]:
            email = _required(row, "Email (khóa nối với CBNV_Import)", "CBNV_BoSung").casefold()
            username = _required(row, "Tên đăng nhập", "CBNV_BoSung")
            codes = [part.strip() for part in _required(row, "Nhóm quyền", "CBNV_BoSung").split(",")]
            if not all(codes) or email in role_by_email:
                raise CommandError(f"Sheet CBNV_BoSung, dòng {row['__row__']}: email hoặc nhóm quyền bị trùng/rỗng.")
            employee_row = employees_by_email.get(email)
            if employee_row is None or _text(employee_row["Tên đăng nhập"]) != username:
                raise CommandError(f"Sheet CBNV_BoSung, dòng {row['__row__']}: không khớp CBNV theo email/username.")
            employee_codes = [part.strip() for part in _text(employee_row["Nhóm quyền"]).split(",")]
            if set(codes) != set(employee_codes):
                raise CommandError(f"Sheet CBNV_BoSung, dòng {row['__row__']}: nhóm quyền không khớp sheet CBNV.")
            role_by_email[email] = codes

        missing_roles = set(employees_by_email) - set(role_by_email)
        extra_roles = set(role_by_email) - set(employees_by_email)
        if missing_roles or extra_roles:
            raise CommandError("Sheet CBNV_BoSung phải có đúng một dòng nhóm quyền cho mỗi CBNV.")

        demo_usernames = {
            _required(row, "Tên đăng nhập", "TaiKhoanDemo")
            for row in records["TaiKhoanDemo"]
        }
        unknown_demo_users = demo_usernames - set(employees_by_username)
        if unknown_demo_users:
            raise CommandError(f"TaiKhoanDemo có username không nằm trong CBNV: {', '.join(sorted(unknown_demo_users))}.")

        for row in records["KiemNhiem"]:
            for column in ("Tên đăng nhập", "Khối kiêm nhiệm", "Đơn vị kiêm nhiệm", "Chức vụ kiêm nhiệm"):
                _required(row, column, "KiemNhiem")
            if _text(row["Tên đăng nhập"]) not in employees_by_username:
                raise CommandError(f"Sheet KiemNhiem, dòng {row['__row__']}: username không có trong CBNV.")

        for sheet_name, columns in {
            "PhongHop": ("Tên phòng", "Vị trí", "Sức chứa", "Trạng thái", "Công khai"),
            "DatPhongHop": ("Phòng", "Người đặt", "Đơn vị", "Tiêu đề", "Ngày lệch", "Bắt đầu", "Kết thúc", "Số người", "Trạng thái"),
            "LoiChucSinhNhat": ("Người nhận", "Người gửi", "Nội dung", "Mẫu thiệp", "Trạng thái", "Ẩn danh"),
        }.items():
            for row in records[sheet_name]:
                for column in columns:
                    _required(row, column, sheet_name)

    def _print_counts(self, records):
        for sheet_name, sheet_records in records.items():
            extra = " (đối chiếu reference data; không ghi bởi lệnh này)" if sheet_name == "DonVi" else ""
            self.stdout.write(f"{sheet_name}: {len(sheet_records)} dòng{extra}")

    def _import_records(self, records, demo_usernames, demo_password):
        UserModel = get_user_model()
        units = list(DonVi.objects.select_related("khoi").order_by("stt_uu_tien", "id"))
        units_by_key = {(unit.khoi.ten, unit.ten): unit for unit in units}
        units_by_name = {}
        for unit in units:
            units_by_name.setdefault(unit.ten, []).append(unit)

        self._verify_unit_reference_data(records["DonVi"], units_by_key)

        positions = {}
        for position in Position.objects.order_by("display_order", "id"):
            positions.setdefault(position.name, position)

        role_by_email = {
            _text(row["Email (khóa nối với CBNV_Import)"]).casefold(): [
                part.strip() for part in _text(row["Nhóm quyền"]).split(",")
            ]
            for row in records["CBNV_BoSung"]
        }
        role_codes = {code for codes in role_by_email.values() for code in codes}
        group_ids = dict(GroupInfo.objects.filter(ma_nhom_quyen__in=role_codes).values_list("ma_nhom_quyen", "id"))
        missing_groups = role_codes - set(group_ids)
        if missing_groups:
            raise CommandError(f"Thiếu nhóm quyền trong DB: {', '.join(sorted(missing_groups))}.")
        groups = {code: Group.objects.get(pk=group_id) for code, group_id in group_ids.items()}

        employee_by_username = {}
        created_employees = 0
        initialized_passwords = 0
        added_group_links = 0
        unusable_password = make_password(None)
        demo_password_hash = make_password(demo_password) if demo_password else None

        for row in records["CBNV"]:
            username = _text(row["Tên đăng nhập"])
            email = _text(row["Email"])
            block_name = _text(row["Khối"])
            unit_name = _text(row["Đơn vị"])
            position_name = _text(row["Chức vụ"])
            unit = units_by_key.get((block_name, unit_name))
            position = positions.get(position_name)
            if unit is None:
                raise CommandError(f"Không tìm thấy đơn vị '{unit_name}' thuộc khối '{block_name}'.")
            if position is None:
                raise CommandError(f"Không tìm thấy chức vụ '{position_name}'.")

            email_owner = UserModel.objects.filter(email__iexact=email).exclude(username=username).first()
            if email_owner:
                raise CommandError(f"Email {email} của {username} đã thuộc tài khoản {email_owner.username}.")

            status = EMPLOYEE_STATUS[_text(row["Trạng thái"])]
            defaults = {
                "ho_ten": _text(row["Họ và tên"]),
                "ngay_sinh": _parse_date(row["Ngày sinh"], "CBNV", row["__row__"], "Ngày sinh"),
                "chuc_vu": position,
                "chuc_danh": _text(row.get("Chức danh")),
                "email": email,
                "sdt": _text(row.get("Số điện thoại")),
                "don_vi": unit,
                "password": demo_password_hash if username in demo_usernames else unusable_password,
                "is_staff": False,
                "is_superuser": False,
                "is_public": _parse_bool(row["Công khai"], "CBNV", row["__row__"], "Công khai"),
                "trang_thai": status,
                "ghi_chu": _text(row.get("Ghi chú")),
            }
            employee = UserModel.objects.filter(username=username).first()
            if employee is None:
                employee = UserModel.objects.create(username=username, **defaults)
                created_employees += 1
            elif employee.email and employee.email.casefold() != email.casefold():
                raise CommandError(f"Username {username} đã tồn tại nhưng email không khớp workbook.")
            elif username in demo_usernames and not employee.has_usable_password():
                employee.set_password(demo_password)
                employee.save(update_fields=["password"])
                initialized_passwords += 1
            employee_by_username[username] = employee

            assigned_groups = set(employee.groups.values_list("pk", flat=True))
            missing_for_user = [groups[code] for code in role_by_email[email.casefold()] if groups[code].pk not in assigned_groups]
            if missing_for_user:
                employee.groups.add(*missing_for_user)
                added_group_links += len(missing_for_user)

        created_assignments = self._import_assignments(records["KiemNhiem"], employee_by_username, units_by_key, positions)
        rooms_by_name, created_rooms = self._import_rooms(records["PhongHop"])
        created_bookings, skipped_bookings = self._import_bookings(
            records["DatPhongHop"], employee_by_username, rooms_by_name, units_by_name
        )
        created_wishes = self._import_birthday_wishes(records["LoiChucSinhNhat"], employee_by_username)

        return {
            "CBNV mới": created_employees,
            "Tài khoản demo được khởi tạo mật khẩu": initialized_passwords,
            "Liên kết nhóm quyền mới": added_group_links,
            "Kiêm nhiệm mới": created_assignments,
            "Phòng họp mới": created_rooms,
            "Lịch đặt mới": created_bookings,
            "Lịch đặt bỏ qua do trùng": skipped_bookings,
            "Lời chúc sinh nhật mới": created_wishes,
        }

    def _verify_unit_reference_data(self, rows, units_by_key):
        for row in rows:
            block_name = _required(row, "Khối", "DonVi")
            unit_name = _required(row, "Tên đơn vị", "DonVi")
            unit = units_by_key.get((block_name, unit_name))
            if unit is None:
                raise CommandError(f"DB thiếu đơn vị '{unit_name}' thuộc khối '{block_name}'.")
            priority = _parse_int(row["STT ưu tiên"], "DonVi", row["__row__"], "STT ưu tiên")
            address = _text(row["Địa chỉ"])
            phone = _text(row["Điện thoại"])
            is_public = _parse_bool(row["Công khai"], "DonVi", row["__row__"], "Công khai")
            if (unit.stt_uu_tien, unit.dia_chi or "", unit.dien_thoai or "", unit.is_public) != (
                priority, address, phone, is_public
            ):
                raise CommandError(f"Thông tin reference data của đơn vị '{unit_name}' chưa khớp workbook; kiểm tra migration organizations.")

    def _import_assignments(self, rows, employees, units_by_key, positions):
        created = 0
        for row in rows:
            username = _text(row["Tên đăng nhập"])
            key = (_text(row["Khối kiêm nhiệm"]), _text(row["Đơn vị kiêm nhiệm"]))
            unit = units_by_key.get(key)
            position = positions.get(_text(row["Chức vụ kiêm nhiệm"]))
            if unit is None or position is None:
                raise CommandError(f"Không phân giải được đơn vị/chức vụ kiêm nhiệm ở dòng {row['__row__']}.")
            title = _text(row.get("Chức danh kiêm nhiệm"))
            assignment, was_created = EmployeeAssignment.objects.get_or_create(
                employee=employees[username],
                don_vi=unit,
                chuc_vu=position,
                defaults={"chuc_danh": title},
            )
            if not was_created and (assignment.chuc_danh or "") != title:
                raise CommandError(f"Kiêm nhiệm của {username} tại '{unit.ten}' đã có chức danh khác.")
            created += int(was_created)
        return created

    def _import_rooms(self, rows):
        rooms = {}
        created = 0
        for row in rows:
            name = _text(row["Tên phòng"])
            status = ROOM_STATUS.get(_text(row["Trạng thái"]))
            if status is None:
                raise CommandError(f"Sheet PhongHop, dòng {row['__row__']}: trạng thái không hợp lệ.")
            room, was_created = MeetingRoom.objects.get_or_create(
                name=name,
                defaults={
                    "location": _text(row["Vị trí"]),
                    "capacity": _parse_int(row["Sức chứa"], "PhongHop", row["__row__"], "Sức chứa"),
                    "amenities": _text(row.get("Tiện ích")),
                    "is_public": _parse_bool(row["Công khai"], "PhongHop", row["__row__"], "Công khai"),
                    "trang_thai": status,
                },
            )
            rooms[name] = room
            created += int(was_created)
        return rooms, created

    def _import_bookings(self, rows, employees, rooms, units_by_name):
        created = 0
        skipped = 0
        today = timezone.localdate()
        for row in rows:
            room_name = _text(row["Phòng"])
            username = _text(row["Người đặt"])
            department_name = _text(row["Đơn vị"])
            room = rooms.get(room_name)
            employee = employees.get(username)
            department_matches = units_by_name.get(department_name, [])
            if room is None or employee is None or len(department_matches) != 1:
                raise CommandError(f"Không phân giải được phòng/người đặt/đơn vị ở sheet DatPhongHop dòng {row['__row__']}.")

            start_time = _parse_time(row["Bắt đầu"], "DatPhongHop", row["__row__"], "Bắt đầu")
            end_time = _parse_time(row["Kết thúc"], "DatPhongHop", row["__row__"], "Kết thúc")
            meeting_date = today + timedelta(days=_parse_int(row["Ngày lệch"], "DatPhongHop", row["__row__"], "Ngày lệch"))
            status = BOOKING_STATUS.get(_text(row["Trạng thái"]))
            if status is None:
                raise CommandError(f"Sheet DatPhongHop, dòng {row['__row__']}: trạng thái không hợp lệ.")
            title = _text(row["Tiêu đề"])
            purpose = _text(row.get("Mục đích"))
            department = department_matches[0]

            stable_match = MeetingBooking.objects.filter(
                room=room,
                organizer=employee,
                department=department,
                title=title,
                start_time=start_time,
                end_time=end_time,
                purpose=purpose,
            ).first()
            if stable_match:
                continue

            if status != MeetingBooking.Status.CANCELLED and MeetingBooking.objects.filter(
                room=room,
                date=meeting_date,
                start_time__lt=end_time,
                end_time__gt=start_time,
            ).exclude(trang_thai=MeetingBooking.Status.CANCELLED).exists():
                skipped += 1
                continue

            MeetingBooking.objects.create(
                room=room,
                organizer=employee,
                department=department,
                title=title,
                date=meeting_date,
                start_time=start_time,
                end_time=end_time,
                attendees_count=_parse_int(row["Số người"], "DatPhongHop", row["__row__"], "Số người"),
                purpose=purpose,
                trang_thai=status,
            )
            created += 1
        return created, skipped

    def _import_birthday_wishes(self, rows, employees):
        created = 0
        for row in rows:
            username = _text(row["Người nhận"])
            status = WISH_STATUS.get(_text(row["Trạng thái"]))
            if username not in employees or status is None:
                raise CommandError(f"Không phân giải được người nhận/trạng thái ở sheet LoiChucSinhNhat dòng {row['__row__']}.")
            _, was_created = BirthdayWish.objects.get_or_create(
                employee=employees[username],
                sender_name=_text(row["Người gửi"]),
                message=_text(row["Nội dung"]),
                defaults={
                    "card_template": _text(row["Mẫu thiệp"]),
                    "is_anonymous": _parse_bool(row["Ẩn danh"], "LoiChucSinhNhat", row["__row__"], "Ẩn danh"),
                    "trang_thai": status,
                },
            )
            created += int(was_created)
        return created
