import os
import tempfile
from datetime import date, timedelta
from io import StringIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from django.utils import timezone

from accounts.models import Employee, EmployeeAssignment, Position
from birthdays.models import BirthdayWish
from meetings.models import MeetingBooking, MeetingRoom
from organizations.models import DonVi


class ImportCommonSampleDataCommandTests(TestCase):
    def setUp(self):
        self.main_unit = DonVi.objects.get(khoi__ten="KHỐI KINH DOANH", ten="Ban dự án lớn")
        self.extra_unit = DonVi.objects.get(
            khoi__ten="KHU VỰC MIỀN TRUNG TÂY NGUYÊN", ten="Bảo Long Phú Yên"
        )

    def create_workbook(self, directory):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "DonVi"
        sheet.append(["STT ưu tiên", "Khối", "Tên đơn vị", "Địa chỉ", "Điện thoại", "Công khai"])
        sheet.append([
            self.main_unit.stt_uu_tien,
            self.main_unit.khoi.ten,
            self.main_unit.ten,
            self.main_unit.dia_chi,
            self.main_unit.dien_thoai,
            "Có" if self.main_unit.is_public else "Không",
        ])

        sheet = workbook.create_sheet("CBNV")
        sheet.append([
            "STT", "Tên đăng nhập", "Họ và tên", "Ngày sinh", "Email", "Số điện thoại", "Đơn vị",
            "Khối", "Chức vụ", "Chức danh", "Nhóm quyền", "Trạng thái", "Công khai", "Ghi chú",
        ])
        sheet.append([
            1, "sample_import_user", "Người dùng mẫu", date(1990, 1, 1), "sample_import_user@example.test",
            "0901234567", self.main_unit.ten, self.main_unit.khoi.ten, "Chuyên viên", "", "CV",
            "Đang làm việc", "Có", "Kiểm thử lệnh nhập",
        ])

        sheet = workbook.create_sheet("CBNV_BoSung")
        sheet.append(["Email (khóa nối với CBNV_Import)", "Tên đăng nhập", "Nhóm quyền"])
        sheet.append(["sample_import_user@example.test", "sample_import_user", "CV"])

        sheet = workbook.create_sheet("KiemNhiem")
        sheet.append([
            "STT", "Tên đăng nhập", "Họ và tên", "Khối chính", "Đơn vị chính", "Chức vụ chính",
            "Khối kiêm nhiệm", "Đơn vị kiêm nhiệm", "Chức vụ kiêm nhiệm", "Chức danh kiêm nhiệm",
        ])
        sheet.append([
            1, "sample_import_user", "Người dùng mẫu", self.main_unit.khoi.ten, self.main_unit.ten,
            "Chuyên viên", self.extra_unit.khoi.ten, self.extra_unit.ten, "Văn thư", "Kiêm nhiệm mẫu",
        ])

        sheet = workbook.create_sheet("TaiKhoanDemo")
        sheet.append(["Tên đăng nhập"])
        sheet.append(["sample_import_user"])

        sheet = workbook.create_sheet("PhongHop")
        sheet.append(["Tên phòng", "Vị trí", "Sức chứa", "Tiện ích", "Công khai", "Trạng thái"])
        sheet.append(["Phòng nhập mẫu test", "Tầng 1", 8, "Bảng trắng", "Có", "Sẵn sàng"])

        sheet = workbook.create_sheet("DatPhongHop")
        sheet.append([
            "STT", "Phòng", "Người đặt", "Đơn vị", "Tiêu đề", "Ngày lệch", "Ngày (mẫu)",
            "Bắt đầu", "Kết thúc", "Số người", "Mục đích", "Trạng thái",
        ])
        sheet.append([
            1, "Phòng nhập mẫu test", "sample_import_user", self.main_unit.ten, "Họp mẫu", 1,
            date(2026, 10, 30), "09:00", "10:00", 2, "Nội dung mẫu", "Đã xác nhận",
        ])

        sheet = workbook.create_sheet("LoiChucSinhNhat")
        sheet.append(["STT", "Người nhận", "Họ tên người nhận", "Người gửi", "Nội dung", "Mẫu thiệp", "Ẩn danh", "Trạng thái"])
        sheet.append([1, "sample_import_user", "Người dùng mẫu", "Người gửi mẫu", "Chúc mừng sinh nhật", "GOLDEN", "Không", "Đã duyệt"])

        path = Path(directory) / "sample.xlsx"
        workbook.save(path)
        workbook.close()
        return path

    def test_default_mode_validates_only(self):
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = self.create_workbook(directory)
            call_command("import_common_sample_data", file=str(workbook_path), stdout=StringIO())

        self.assertFalse(Employee.objects.filter(username="sample_import_user").exists())

    def test_apply_imports_idempotently_and_uses_env_password_for_demo_user(self):
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = self.create_workbook(directory)
            demo_password = " temporary-test-password "
            with patch.dict(os.environ, {"DEMO_PASSWORD": demo_password}):
                call_command("import_common_sample_data", file=str(workbook_path), apply=True, stdout=StringIO())

                employee = Employee.objects.get(username="sample_import_user")
                self.assertTrue(employee.check_password(demo_password))
                self.assertFalse(employee.is_superuser)
                self.assertTrue(employee.groups.filter(name="Chuyên viên").exists())
                self.assertTrue(
                    EmployeeAssignment.objects.filter(
                        employee=employee,
                        don_vi=self.extra_unit,
                        chuc_vu__name="Văn thư",
                        chuc_danh="Kiêm nhiệm mẫu",
                    ).exists()
                )
                room = MeetingRoom.objects.get(name="Phòng nhập mẫu test")
                booking = MeetingBooking.objects.get(room=room, organizer=employee, title="Họp mẫu")
                self.assertEqual(booking.date, timezone.localdate() + timedelta(days=1))
                self.assertTrue(BirthdayWish.objects.filter(employee=employee, message="Chúc mừng sinh nhật").exists())

                employee.set_password("existing-password")
                employee.save(update_fields=["password"])
                call_command("import_common_sample_data", file=str(workbook_path), apply=True, stdout=StringIO())

        self.assertEqual(Employee.objects.filter(username="sample_import_user").count(), 1)
        self.assertEqual(EmployeeAssignment.objects.filter(employee__username="sample_import_user").count(), 1)
        self.assertEqual(MeetingBooking.objects.filter(organizer__username="sample_import_user").count(), 1)
        self.assertEqual(BirthdayWish.objects.filter(employee__username="sample_import_user").count(), 1)
        self.assertTrue(Employee.objects.get(username="sample_import_user").check_password("existing-password"))

    def test_apply_requires_demo_password_before_writing(self):
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = self.create_workbook(directory)
            with patch.dict(os.environ, {"DEMO_PASSWORD": ""}):
                with self.assertRaises(CommandError):
                    call_command("import_common_sample_data", file=str(workbook_path), apply=True, stdout=StringIO())

        self.assertFalse(Employee.objects.filter(username="sample_import_user").exists())
