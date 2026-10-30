"""
Hỗ trợ xuất tệp mẫu và nhập dữ liệu Cán bộ Nhân viên (CBNV) bằng Excel (.xlsx).
Tác giả: Lê Duy Khánh (leduykhanhdevs).
"""
import re
import secrets
import unicodedata
from io import BytesIO
from datetime import datetime, date
from typing import Callable
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.contrib.auth.models import Group
from django.db import transaction
from django.contrib.auth import get_user_model

from organizations.models import DonVi, KhoiKhuVuc
from .models import Employee, EmployeeAssignment, Position, vietnam_phone_regex
from .serializers import get_default_groups_for_position


def generate_cbnv_template() -> BytesIO:
    """
    Sinh tệp Excel mẫu để nhập danh sách CBNV.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh Sach CBNV"

    # Header styling
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    hint_font = Font(name="Arial", size=9, italic=True, color="475569")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    headers = [
        "Họ và tên (*)",
        "Ngày sinh (*)",
        "Email (*)",
        "Số điện thoại",
        "Tên Khối / Khu vực (*)",
        "Tên Đơn vị / Phòng ban (*)",
        "Tên Chức vụ (*)",
        "Chức danh",
        "Trạng thái (Đang làm việc / Nghỉ việc)",
        "Ghi chú",
    ]

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = thin_border

    # Sample rows
    sample_rows = [
        [
            "Nguyễn Văn An",
            "1990-05-15",
            "an.nguyen@baolong.vn",
            "0912345678",
            "KHU VỰC THÀNH PHỐ HỒ CHÍ MINH",
            "Bảo Long Sài Gòn",
            "Trưởng phòng",
            "Trưởng phòng Kinh doanh",
            "Đang làm việc",
            "Nhân sự tuyển dụng đợt 1",
        ],
        [
            "Trần Thị Bích",
            "1994-11-20",
            "bich.tran@baolong.vn",
            "0987654321",
            "KHỐI HỖ TRỢ",
            "Ban Pháp chế",
            "Chuyên viên",
            "Chuyên viên Pháp lý",
            "Đang làm việc",
            "Nhân sự chính thức",
        ],
    ]

    # Hints row
    hints = [
        "(Ví dụ: Nguyễn Văn A)",
        "(Định dạng YYYY-MM-DD)",
        "(Ví dụ: an.nguyen@baolong.vn)",
        "(Ví dụ: 0912345678)",
        "(Tên khối/khu vực đã có trong hệ thống)",
        "(Tên đơn vị thuộc đúng khối/khu vực)",
        "(Tên chức vụ)",
        "(Chức danh nội bộ)",
        "(Đang làm việc / Nghỉ việc)",
        "(Ghi chú tự do)",
    ]
    ws.append(hints)
    for col_idx in range(1, len(hints) + 1):
        cell = ws.cell(row=2, column=col_idx)
        cell.font = hint_font
        cell.alignment = align_center
        cell.border = thin_border

    for row_data in sample_rows:
        ws.append(row_data)

    # Set column widths
    col_widths = {
        "A": 26,  # Họ và tên
        "B": 16,  # Ngày sinh
        "C": 28,  # Email
        "D": 18,  # SĐT
        "E": 34,  # Khối / khu vực
        "F": 30,  # Đơn vị
        "G": 22,  # Chức vụ
        "H": 24,  # Chức danh
        "I": 25,  # Trạng thái
        "J": 30,  # Ghi chú
    }
    for col_letter, w in col_widths.items():
        ws.column_dimensions[col_letter].width = w

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _parse_date(val) -> date | None:
    if isinstance(val, (datetime, date)):
        return val.date() if isinstance(val, datetime) else val
    if not val:
        return None
    val_str = str(val).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(val_str, fmt).date()
        except ValueError:
            continue
    return None



def _read_rows_from_excel(file_obj) -> list[tuple]:
    """
    Đọc dữ liệu từ file Excel định dạng .xlsx qua openpyxl.
    """
    wb = openpyxl.load_workbook(file_obj, data_only=True)
    try:
        return list(wb.active.iter_rows(values_only=True))
    finally:
        wb.close()


def _normalize_excel_header(value) -> str:
    """Chuẩn hóa tiêu đề để nhận cả file mẫu và file CBNV đã xuất từ hệ thống."""
    if value is None:
        return ""
    text = str(value).strip().casefold().replace("đ", "d")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]", "", text)


def _map_excel_headers(header_row: tuple) -> dict[str, int]:
    """Map các tên cột phổ biến về trường import chuẩn; bỏ qua cột không dùng."""
    aliases = {
        "ho_ten": {"hovaten", "hoten"},
        "ngay_sinh": {"ngaysinh"},
        "email": {"email", "emailcongvu"},
        "sdt": {"sodienthoai", "sdt"},
        "khoi": {"khoi", "tenkhoi", "khoikhuvuc", "tenkhoikhuvuc"},
        "don_vi": {"donvi", "tendonvi", "tendonviphongban"},
        "chuc_vu": {"chucvu", "tenchucvu"},
        "chuc_danh": {"chucdanh"},
        "trang_thai": {"trangthai"},
        "ghi_chu": {"ghichu"},
    }
    alias_to_field = {
        alias: field for field, field_aliases in aliases.items() for alias in field_aliases
    }
    mapped = {}
    for index, value in enumerate(header_row):
        key = _normalize_excel_header(value)
        field = alias_to_field.get(key)
        if field and field not in mapped:
            mapped[field] = index
        elif key.startswith("trangthai") and "trang_thai" not in mapped:
            mapped["trang_thai"] = index
    return mapped


def import_cbnv_from_excel(
    file_obj,
    progress_callback: Callable[[str, int, int], None] | None = None,
    completion_callback: Callable[[int], None] | None = None,
) -> dict:
    """
    Đọc và kiểm tra tính hợp lệ từng dòng của tệp Excel nhập danh sách CBNV (.xlsx).
    Trả về dict gồm:
      - success: bool
      - message: str
      - created_count: int
      - errors: list[dict(row=int, message=str)]
    """
    try:
        rows = _read_rows_from_excel(file_obj)
    except Exception as e:
        return {
            "success": False,
            "message": "employees.import.errors.invalidFile",
            "errors": [{
                "row": 0,
                "field": "file",
                "code": "employees.import.errors.invalidFile",
                "value": "",
                "message": f"Không thể đọc tệp Excel. Lỗi: {str(e)}",
            }],
        }
    if len(rows) <= 1:
        return {
            "success": False,
            "message": "employees.import.errors.emptyFile",
            "errors": [{
                "row": 0,
                "field": "file",
                "code": "employees.import.errors.emptyFile",
                "value": "",
                "message": "Tệp Excel trống hoặc không có dữ liệu ngoài dòng tiêu đề.",
            }],
        }

    # Ưu tiên ánh xạ theo tiêu đề để hỗ trợ cả mẫu mới có Khối và file CBNV đã xuất.
    # Giữ fallback theo vị trí cho các file cũ không có tiêu đề nhận diện được.
    column_indices = _map_excel_headers(rows[0])
    required_header_fields = {"ho_ten", "ngay_sinh", "email", "sdt", "don_vi", "chuc_vu"}
    has_header_mapping = required_header_fields.issubset(column_indices)
    has_block_column = has_header_mapping and "khoi" in column_indices
    offset = 0
    if not has_header_mapping:
        first_row = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
        offset = 1 if (len(first_row) > 0 and ("stt" in first_row[0] or "số tt" in first_row[0] or "số thứ tự" in first_row[0])) else 0
        column_indices = {
            "ho_ten": 0 + offset,
            "ngay_sinh": 1 + offset,
            "email": 2 + offset,
            "sdt": 3 + offset,
            "don_vi": 4 + offset,
            "chuc_vu": 5 + offset,
            "chuc_danh": 6 + offset,
            "trang_thai": 7 + offset,
            "ghi_chu": 8 + offset,
        }

    # Đọc trước các danh mục và email hiện có để tránh lặp nhiều truy vấn theo từng dòng.
    # Giữ toàn bộ đơn vị trùng tên để phân biệt bằng Khối / Khu vực khi file có cột này.
    def _comparison_key(value):
        return str(value or "").upper()

    UserModel = get_user_model()
    existing_email_keys = {
        _comparison_key(value)
        for value in UserModel.objects.exclude(email__isnull=True)
        .exclude(email="")
        .values_list("email", flat=True)
    }
    units_by_name = {}
    units_by_name_and_block = {}
    for unit in DonVi.objects.select_related("khoi").all():
        name_key = _comparison_key(unit.ten)
        block_key = _comparison_key(unit.khoi.ten if unit.khoi_id else "")
        units_by_name.setdefault(name_key, []).append(unit)
        units_by_name_and_block.setdefault((name_key, block_key), []).append(unit)

    positions_by_name = {}
    for position in Position.objects.all():
        positions_by_name.setdefault(_comparison_key(position.name), position)

    errors = []
    valid_records = []
    seen_emails_in_file = set()

    def add_error(row: int, field: str, code: str, value="", message=""):
        errors.append({
            "row": row,
            "field": field,
            "code": code,
            "value": "" if value is None else str(value),
            "message": message,
        })

    total_rows = len(rows) - 1
    for idx, row in enumerate(rows[1:], start=2):
        processed_rows = idx - 1
        if progress_callback and (processed_rows % 10 == 0 or processed_rows == total_rows):
            progress_callback("validation", processed_rows, total_rows)
        if not any(row):
            continue  # Bỏ qua dòng trống hoàn toàn

        # Bỏ qua dòng hướng dẫn / gợi ý nếu ô bắt đầu bằng dấu ngoặc đơn '('
        first_cell = str(row[0]).strip() if len(row) > 0 and row[0] is not None else ""
        second_cell = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        if first_cell.startswith("(") or second_cell.startswith("("):
            continue

        def _get_val(field):
            target = column_indices.get(field)
            if target is None:
                return None
            if target < len(row) and row[target] is not None:
                return row[target]
            return None

        ho_ten = str(_get_val("ho_ten")).strip() if _get_val("ho_ten") is not None else ""
        raw_birthday = _get_val("ngay_sinh")
        raw_email = str(_get_val("email")).strip().lower() if _get_val("email") is not None else ""
        sdt = str(_get_val("sdt")).strip() if _get_val("sdt") is not None else ""
        khoi_name = str(_get_val("khoi")).strip() if _get_val("khoi") is not None else ""
        don_vi_name = str(_get_val("don_vi")).strip() if _get_val("don_vi") is not None else ""
        chuc_vu_name = str(_get_val("chuc_vu")).strip() if _get_val("chuc_vu") is not None else ""
        chuc_danh = str(_get_val("chuc_danh")).strip() if _get_val("chuc_danh") is not None else ""
        raw_trang_thai = str(_get_val("trang_thai")).strip() if _get_val("trang_thai") is not None else "Đang làm việc"
        ghi_chu = str(_get_val("ghi_chu")).strip() if _get_val("ghi_chu") is not None else ""

        # 1. Validate Họ và tên
        if not ho_ten:
            add_error(idx, "ho_ten", "employees.validation.nameRequired", "", "Họ và tên không được để trống.")

        # 2. Validate Ngày sinh
        birthday = _parse_date(raw_birthday)
        if not birthday:
            add_error(
                idx,
                "ngay_sinh",
                "employees.import.errors.invalidBirthday",
                str(raw_birthday or ""),
                f"Ngày sinh '{raw_birthday}' không đúng định dạng (hỗ trợ YYYY-MM-DD hoặc DD/MM/YYYY).",
            )
        else:
            today = date.today()
            if birthday > today:
                add_error(
                    idx,
                    "ngay_sinh",
                    "employees.messages.birthdayInFuture",
                    str(raw_birthday or ""),
                    "Ngày sinh không thể ở trong tương lai.",
                )
            else:
                age = today.year - birthday.year - ((today.month, today.day) < (birthday.month, birthday.day))
                if age < 18 or age > 100:
                    add_error(
                        idx,
                        "ngay_sinh",
                        "employees.messages.birthdayOutOfRange",
                        str(raw_birthday or ""),
                        "Độ tuổi của cán bộ nhân viên phải từ 18 đến 100 tuổi.",
                    )

        # 3. Validate Email
        if not raw_email:
            add_error(idx, "email", "employees.validation.emailRequired", "", "Email không được để trống.")
        elif not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", raw_email):
            add_error(
                idx,
                "email",
                "employees.import.errors.invalidEmail",
                raw_email,
                f"Email '{raw_email}' không đúng định dạng.",
            )
        elif raw_email in seen_emails_in_file:
            add_error(
                idx,
                "email",
                "employees.import.errors.duplicateEmailInFile",
                raw_email,
                f"Email '{raw_email}' bị trùng lặp trong chính tệp Excel tải lên.",
            )
        elif _comparison_key(raw_email) in existing_email_keys:
            add_error(
                idx,
                "email",
                "employees.messages.emailExists",
                raw_email,
                f"Email '{raw_email}' đã tồn tại trên hệ thống.",
            )
        else:
            seen_emails_in_file.add(raw_email)

        # 4. Validate Số điện thoại
        if sdt:
            sdt_clean = sdt.replace(" ", "").replace(".", "")
            if not re.match(r"^(0|\+84)[1-9][0-9]{7,9}$", sdt_clean):
                add_error(
                    idx,
                    "sdt",
                    "employees.validation.phoneInvalid",
                    sdt,
                    f"Số điện thoại '{sdt}' không hợp lệ theo chuẩn Việt Nam.",
                )
        else:
            sdt_clean = ""

        # 5. Validate Đơn vị / Phòng ban
        don_vi_instance = None
        if has_block_column and not khoi_name:
            add_error(
                idx,
                "khoi",
                "employees.import.errors.blockRequired",
                "",
                "Tên Khối / Khu vực không được để trống.",
            )
        if not don_vi_name:
            add_error(idx, "don_vi", "employees.import.errors.unitRequired", "", "Tên Đơn vị / Phòng ban không được để trống.")
        elif has_block_column and not khoi_name:
            # Đã báo lỗi thiếu Khối; không tra cứu không phân biệt Khối để tránh gán nhầm đơn vị.
            pass
        else:
            unit_key = _comparison_key(don_vi_name)
            matching_units = (
                units_by_name_and_block.get((unit_key, _comparison_key(khoi_name)), [])
                if has_block_column
                else units_by_name.get(unit_key, [])
            )
            unit_count = len(matching_units)
            if unit_count == 0:
                scope = f" trong Khối / Khu vực '{khoi_name}'" if has_block_column else ""
                add_error(
                    idx,
                    "don_vi",
                    "employees.messages.unitNotFound",
                    don_vi_name,
                    f"Không tìm thấy Đơn vị có tên '{don_vi_name}'{scope} trong hệ thống.",
                )
            elif unit_count > 1:
                if has_block_column:
                    ambiguous_message = (
                        f"Tên Đơn vị '{don_vi_name}' vẫn trùng trong Khối / Khu vực "
                        f"'{khoi_name}'. Vui lòng kiểm tra danh mục đơn vị."
                    )
                else:
                    ambiguous_message = (
                        f"Tên Đơn vị '{don_vi_name}' trùng lặp giữa nhiều Khối/Khu vực. "
                        "Vui lòng dùng mẫu mới và điền cột Khối / Khu vực."
                    )
                add_error(
                    idx,
                    "don_vi",
                    "employees.messages.unitAmbiguous",
                    don_vi_name,
                    ambiguous_message,
                )
            else:
                don_vi_instance = matching_units[0]

        # 6. Validate Chức vụ
        chuc_vu_instance = None
        if not chuc_vu_name:
            add_error(idx, "chuc_vu", "employees.import.errors.positionRequired", "", "Tên Chức vụ không được để trống.")
        else:
            chuc_vu_instance = positions_by_name.get(_comparison_key(chuc_vu_name))
            if not chuc_vu_instance:
                add_error(
                    idx,
                    "chuc_vu",
                    "employees.messages.positionNotFound",
                    chuc_vu_name,
                    f"Không tìm thấy Chức vụ có tên '{chuc_vu_name}' trong hệ thống.",
                )

        # 7. Validate Trạng thái
        trang_thai = Employee.Status.LAM_VIEC
        if "nghỉ" in raw_trang_thai.lower() or "nghi" in raw_trang_thai.lower():
            trang_thai = Employee.Status.NGHI_VIEC

        valid_records.append({
            "ho_ten": ho_ten,
            "ngay_sinh": birthday,
            "email": raw_email,
            "sdt": sdt_clean,
            "don_vi": don_vi_instance,
            "chuc_vu": chuc_vu_instance,
            "chuc_danh": chuc_danh,
            "trang_thai": trang_thai,
            "ghi_chu": ghi_chu,
        })

    if progress_callback and total_rows:
        progress_callback("validation", total_rows, total_rows)

    if errors:
        return {
            "success": False,
            "errors": errors,
            "message": "employees.messages.importHasErrors",
        }

    # Nếu toàn bộ dữ liệu hợp lệ, tiến hành lưu nguyên tử (atomic) vào CSDL
    created_count = 0
    existing_username_keys = {
        _comparison_key(value)
        for value in UserModel.objects.values_list("username", flat=True)
    }
    default_groups_by_position = {}
    total_to_create = len(valid_records)
    batch_size = 100
    with transaction.atomic():
        groups_through = UserModel.groups.through
        owner_field = next(
            field for field in groups_through._meta.fields
            if field.is_relation and field.related_model is UserModel
        )
        group_field = next(
            field for field in groups_through._meta.fields
            if field.is_relation and field.related_model is Group
        )

        for batch_start in range(0, total_to_create, batch_size):
            batch = valid_records[batch_start:batch_start + batch_size]
            imported_users = []
            groups_for_batch = []
            for item in batch:
                email = item["email"]
                local_part = email.split("@")[0]
                candidate = local_part
                suffix = 1
                while _comparison_key(candidate) in existing_username_keys:
                    candidate = f"{local_part}{suffix}"
                    suffix += 1
                existing_username_keys.add(_comparison_key(candidate))

                imported_user = UserModel(
                    username=candidate,
                    email=UserModel.objects.normalize_email(email),
                    ho_ten=item["ho_ten"],
                    ngay_sinh=item["ngay_sinh"],
                    don_vi=item["don_vi"],
                    chuc_vu=item["chuc_vu"],
                    chuc_danh=item["chuc_danh"],
                    sdt=item["sdt"],
                    trang_thai=item["trang_thai"],
                    is_active=(item["trang_thai"] != Employee.Status.NGHI_VIEC),
                    ghi_chu=item["ghi_chu"],
                )
                imported_user.set_password(f"ITC@{secrets.token_hex(4)}")
                imported_users.append(imported_user)

                position_id = item["chuc_vu"].pk
                if position_id not in default_groups_by_position:
                    default_groups_by_position[position_id] = get_default_groups_for_position(item["chuc_vu"])
                groups_for_batch.append(default_groups_by_position[position_id])

            # `trang_thai` is a PostgreSQL ENUM in release_DDL.sql. Django's
            # PostgreSQL bulk INSERT uses UNNEST(varchar[]) and cannot assign
            # that array to employee_work_status, so keep model inserts typed
            # through the regular save path while batching the M2M group rows.
            for imported_user in imported_users:
                imported_user.save(force_insert=True)
            group_rows = [
                groups_through(**{
                    owner_field.attname: imported_user.pk,
                    group_field.attname: group.pk,
                })
                for imported_user, default_groups in zip(imported_users, groups_for_batch)
                for group in default_groups
            ]
            if group_rows:
                groups_through.objects.bulk_create(group_rows, batch_size=500)

            created_count += len(imported_users)
            if progress_callback:
                progress_callback("saving", created_count, total_to_create)

        # Commit trạng thái hoàn tất cùng transaction tạo CBNV để job không thể bị chạy lại
        # nếu tiến trình dừng ngay sau khi dữ liệu đã được commit.
        if completion_callback:
            completion_callback(created_count)

    if progress_callback and total_to_create:
        progress_callback("saving", total_to_create, total_to_create)

    return {
        "success": True,
        "message": "employees.messages.importSuccess",
        "created_count": created_count,
        "errors": [],
    }


def _write_cell(cell, value=None, font=None, fill=None, alignment=None, border=None):
    from openpyxl.cell.cell import MergedCell

    if not isinstance(cell, MergedCell) and value is not None:
        cell.value = value
        if isinstance(value, str):
            cell.data_type = "s"
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if alignment:
        cell.alignment = alignment
    if border:
        cell.border = border


def export_cbnv_to_excel(queryset, user) -> BytesIO:
    """
    Xuất danh bạ Cán bộ Nhân viên theo phân cấp Khối -> Đơn vị -> Nhân sự (Issue #24).
    Tuân thủ quy chuẩn của thầy & phản hồi trưởng nhóm:
    - Khối sắp xếp theo thứ tự hiển thị 1 đến 16 (thu_tu_hien_thi).
    - Đơn vị sắp xếp theo STT ưu tiên (stt_uu_tien).
    - Nhân sự trong đơn vị sắp xếp theo thứ tự cấp bậc chức vụ (Position.display_order), sau đó theo họ tên.
    - STT đếm lại từ 1 ở mỗi đơn vị.
    - Nhân sự kiêm nhiệm (EmployeeAssignment) xuất hiện dưới từng đơn vị họ đảm nhiệm.
    - 6 cột: STT, Họ và tên, Chức vụ, Chức danh, Điện thoại, Email.
    - Lọc CBNV đang làm việc (trang_thai = 'LAM_VIEC', is_active = True).
    - CBNV Nội bộ (is_public = False) chỉ xuất hiện với người có quyền view_full_cbnv hoặc superuser.
    - Chống Formula Injection: ép kiểu chuỗi an toàn cell.data_type = 's'.
    """
    can_view_full = bool(
        user
        and (
            user.is_superuser
            or user.has_perm("authorization.view_full_cbnv")
            or user.has_perm("accounts.view_full_cbnv")
            or user.has_perm("view_full_cbnv")
        )
    )
    can_view_internal_units = bool(
        user
        and (
            user.is_superuser
            or user.has_perm("organizations.change_donvi")
            or user.has_perm("change_donvi")
        )
    )

    # Lọc chỉ CBNV đang làm việc, loại bỏ tài khoản superuser/quản trị hệ thống mẫu
    active_employees = (
        queryset.filter(
            trang_thai=Employee.Status.LAM_VIEC,
            is_active=True,
        )
        .exclude(is_superuser=True)
        .exclude(email__iexact="admin@baolong.vn")
        .exclude(username="admin")
    )
    if not can_view_full:
        active_employees = active_employees.filter(is_public=True)

    active_employees = (
        active_employees.select_related("don_vi__khoi", "chuc_vu")
        .prefetch_related("assignments__don_vi__khoi", "assignments__chuc_vu")
        .all()
    )

    # Gom danh sách nhân sự theo don_vi_id
    unit_persons = {}
    active_emp_ids = set()

    for emp in active_employees:
        active_emp_ids.add(emp.id)
        if emp.don_vi_id:
            if not can_view_internal_units and emp.don_vi and not emp.don_vi.is_public:
                continue
            disp_order = emp.chuc_vu.display_order if emp.chuc_vu else 999999
            unit_persons.setdefault(emp.don_vi_id, []).append({
                "ho_ten": emp.ho_ten,
                "chuc_vu": emp.chuc_vu.name if emp.chuc_vu else "",
                "chuc_danh": emp.chuc_danh or "",
                "sdt": emp.sdt or "",
                "email": emp.email or "",
                "display_order": disp_order,
            })

    # Phân công kiêm nhiệm của các nhân sự đang làm việc
    if active_emp_ids:
        assignments_qs = (
            EmployeeAssignment.objects.filter(employee_id__in=active_emp_ids)
            .select_related("employee", "don_vi__khoi", "chuc_vu")
        )
        if not can_view_internal_units:
            assignments_qs = assignments_qs.filter(don_vi__is_public=True)

        for assign in assignments_qs.all():
            if assign.don_vi_id:
                disp_order = assign.chuc_vu.display_order if assign.chuc_vu else 999999
                raw_chuc_danh = (assign.chuc_danh or "").strip()
                if not raw_chuc_danh:
                    chuc_danh = "Kiêm nhiệm"
                elif "kiêm nhiệm" not in raw_chuc_danh.lower():
                    chuc_danh = f"{raw_chuc_danh} (Kiêm nhiệm)"
                else:
                    chuc_danh = raw_chuc_danh

                unit_persons.setdefault(assign.don_vi_id, []).append({
                    "ho_ten": assign.employee.ho_ten,
                    "chuc_vu": assign.chuc_vu.name if assign.chuc_vu else "",
                    "chuc_danh": chuc_danh,
                    "sdt": assign.employee.sdt or "",
                    "email": assign.employee.email or "",
                    "display_order": disp_order,
                })

    all_don_vi_ids = list(unit_persons.keys())
    don_vis_qs = DonVi.objects.filter(id__in=all_don_vi_ids)
    if not can_view_internal_units:
        don_vis_qs = don_vis_qs.filter(is_public=True)
    don_vis = don_vis_qs.select_related("khoi").order_by("stt_uu_tien", "id")

    khoi_units_map = {}
    don_vis_without_khoi = []
    khoi_map = {}

    for dv in don_vis:
        if dv.khoi:
            khoi_map[dv.khoi.id] = dv.khoi
            khoi_units_map.setdefault(dv.khoi.id, []).append(dv)
        else:
            don_vis_without_khoi.append(dv)

    sorted_khois = sorted(khoi_map.values(), key=lambda k: (k.thu_tu_hien_thi, k.id))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh Ba CBNV"

    title_font = Font(name="Arial", size=13, bold=True, color="1E3A8A")
    sub_font_bold = Font(name="Arial", size=9, bold=True, color="1E293B")
    sub_font_italic = Font(name="Arial", size=9, italic=True, color="475569")
    header_font = Font(name="Arial", size=10, bold=True, color="000000")
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    khoi_font = Font(name="Arial", size=11, bold=True, color="1E3A8A")
    khoi_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    unit_font = Font(name="Arial", size=10, bold=True, color="000000")
    unit_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    body_font = Font(name="Arial", size=10, color="000000")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    # 1. Header Banner
    ws.merge_cells("A1:F1")
    _write_cell(ws["A1"], "Bảo Long Insurance — Hệ Thống Quản Lý Văn Bản & Điều Hành", font=title_font, alignment=align_center)
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:F2")
    _write_cell(ws["A2"], "Danh bạ cán bộ nhân viên", font=sub_font_bold, alignment=align_center)
    ws.row_dimensions[2].height = 20

    ws.merge_cells("A3:F3")
    _write_cell(ws["A3"], f"Ngày xuất: {datetime.now().strftime('%d/%m/%Y')}", font=sub_font_italic, alignment=align_center)
    ws.row_dimensions[3].height = 18

    ws.row_dimensions[4].height = 10

    # 2. Table Column Headers
    headers = ["STT", "Họ và tên", "Chức vụ", "Chức danh", "Điện thoại", "Email"]
    ws.row_dimensions[5].height = 26
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col_idx)
        _write_cell(cell, h, font=header_font, fill=header_fill, alignment=align_center, border=thin_border)

    current_row = 6

    def render_unit(dv):
        nonlocal current_row
        parts = [dv.ten.upper()]
        if dv.dia_chi:
            parts.append(f"ĐỊA CHỈ: {dv.dia_chi}")
        if dv.dien_thoai:
            parts.append(f"ĐT: {dv.dien_thoai}")
        if dv.email:
            parts.append(f"EMAIL: {dv.email}")
        unit_text = "   ".join(parts)

        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=6)
        for c in range(1, 7):
            _write_cell(ws.cell(row=current_row, column=c), unit_text if c == 1 else None, font=unit_font, fill=unit_fill, alignment=align_left, border=thin_border)
        ws.row_dimensions[current_row].height = 22
        current_row += 1

        persons = unit_persons.get(dv.id, [])
        persons.sort(key=lambda p: (p["display_order"], p["ho_ten"]))

        for p_idx, person in enumerate(persons, 1):
            row_fill = zebra_fill if p_idx % 2 == 0 else white_fill
            row_vals = [
                p_idx,
                person["ho_ten"],
                person["chuc_vu"],
                person["chuc_danh"],
                person["sdt"],
                person["email"],
            ]
            ws.row_dimensions[current_row].height = 20
            for col_idx, val in enumerate(row_vals, 1):
                cell = ws.cell(row=current_row, column=col_idx)
                align = align_center if col_idx in (1, 5) else align_left
                _write_cell(cell, val, font=body_font, fill=row_fill, alignment=align, border=thin_border)
            current_row += 1

    for khoi in sorted_khois:
        units_in_khoi = khoi_units_map.get(khoi.id, [])
        if not units_in_khoi:
            continue
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=6)
        for c in range(1, 7):
            _write_cell(ws.cell(row=current_row, column=c), khoi.ten.upper() if c == 1 else None, font=khoi_font, fill=khoi_fill, alignment=align_left, border=thin_border)
        ws.row_dimensions[current_row].height = 24
        current_row += 1

        for dv in units_in_khoi:
            render_unit(dv)

    if don_vis_without_khoi:
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=6)
        for c in range(1, 7):
            _write_cell(ws.cell(row=current_row, column=c), "ĐƠN VỊ KHÁC" if c == 1 else None, font=khoi_font, fill=khoi_fill, alignment=align_left, border=thin_border)
        ws.row_dimensions[current_row].height = 24
        current_row += 1
        for dv in don_vis_without_khoi:
            render_unit(dv)

    col_min_widths = {1: 8, 2: 24, 3: 20, 4: 24, 5: 16, 6: 28}
    for col_idx, min_w in col_min_widths.items():
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        max_len = min_w
        for r in range(5, current_row):
            if ws.row_dimensions[r].height == 20:
                c_val = str(ws.cell(row=r, column=col_idx).value or "")
                if len(c_val) > max_len:
                    max_len = len(c_val)
        ws.column_dimensions[col_letter].width = max(max_len + 3, min_w)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
