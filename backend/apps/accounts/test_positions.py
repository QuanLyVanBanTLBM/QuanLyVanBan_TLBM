"""
Test Danh mục Chức vụ (Issue #20): model, CRUD API, quyền RBAC, STT tự tăng, import/template Excel.
"""
from io import BytesIO
from unittest import mock

import openpyxl
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError
from django.test import TestCase
from rest_framework import status
from rest_framework.test import APIClient

from accounts.excel_position import HEADERS, MAX_ROWS
from accounts.models import Employee, Position
from organizations.models import DonVi, KhoiKhuVuc

LIST_URL = "/api/positions/"
TEMPLATE_URL = "/api/positions/import-template/"
IMPORT_URL = "/api/positions/import/"


def detail_url(pk):
    return f"{LIST_URL}{pk}/"


def make_xlsx(rows, header=None, name="positions.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(header or HEADERS)
    for row in rows:
        ws.append(row)
    buffer = BytesIO()
    wb.save(buffer)
    return SimpleUploadedFile(
        name,
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


class PositionBaseTestCase(TestCase):
    def setUp(self):
        User = get_user_model()

        def make_user(username, *codenames):
            user = User.objects.create_user(username=username, password="Test@12345")
            if codenames:
                user.user_permissions.add(*Permission.objects.filter(codename__in=codenames))
            return user

        self.viewer = make_user("pos_viewer", "view_chucvu")
        self.manager = make_user(
            "pos_manager", "view_chucvu", "add_chucvu", "change_chucvu", "delete_chucvu", "import_chucvu"
        )
        self.no_perm = make_user("pos_noperm")
        self.root = User.objects.create_superuser(username="pos_root", password="Test@12345", email="r@example.com")

        self.client_for = {}
        for user in (self.viewer, self.manager, self.no_perm, self.root):
            client = APIClient()
            client.force_authenticate(user=user)
            self.client_for[user.username] = client
        self.anon = APIClient()

        # DB test dựng bằng migration có 25 chức vụ chuẩn (accounts/0009); các ca kiểm danh sách, phân trang,
        # import chỉ cần chức vụ đang được tài khoản dùng ("Quản trị viên" của admin), bỏ các chức vụ còn lại.
        Position.objects.exclude(pk__in=Employee.objects.values("chuc_vu_id")).delete()
        self.baseline_position_ids = list(Position.objects.values_list("pk", flat=True))
        self.baseline_position_count = len(self.baseline_position_ids)
        self.baseline_public_count = Position.objects.filter(is_public=True).count()

    def others(self):
        return Position.objects.exclude(pk__in=self.baseline_position_ids)

    def next_order(self, offset=0):
        return Position.get_next_display_order() + offset


class PositionModelTests(PositionBaseTestCase):
    def test_display_order_must_be_unique(self):
        duplicate_order = self.next_order()
        Position.objects.create(display_order=duplicate_order, name="A")
        with self.assertRaises(IntegrityError):
            Position.objects.create(display_order=duplicate_order, name="B")

    def test_next_display_order_is_max_plus_one(self):
        first_available = self.next_order()
        self.assertEqual(Position.get_next_display_order(), first_available)
        Position.objects.create(display_order=first_available + 4, name="A")
        Position.objects.create(display_order=first_available + 1, name="B")
        self.assertEqual(Position.get_next_display_order(), first_available + 5)


class PositionListTests(PositionBaseTestCase):
    def setUp(self):
        super().setUp()
        self.first_fixture_order = self.next_order()
        Position.objects.create(display_order=self.first_fixture_order + 2, name="Chuyên viên kiểm thử")
        Position.objects.create(display_order=self.first_fixture_order, name="Tổng Giám đốc kiểm thử")
        Position.objects.create(display_order=self.first_fixture_order + 1, name="Trưởng phòng kiểm thử")
        Position.objects.create(
            display_order=self.first_fixture_order + 3, name="Cộng tác viên kiểm thử", is_public=False
        )

    def test_unauthenticated_401(self):
        self.assertEqual(self.anon.get(LIST_URL).status_code, status.HTTP_401_UNAUTHORIZED)

    def test_without_view_permission_403(self):
        res = self.client_for["pos_noperm"].get(LIST_URL)
        self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(res.json()["detail"], "positions.messages.permissionDenied")

    def test_list_sorted_by_stt_with_issue_field_names(self):
        res = self.client_for["pos_manager"].get(LIST_URL, {"all": "true"})
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        results = [
            p for p in res.json()
            if p["name"].endswith(" kiểm thử") and p["name"] in {
                "Tổng Giám đốc kiểm thử", "Trưởng phòng kiểm thử",
                "Chuyên viên kiểm thử", "Cộng tác viên kiểm thử",
            }
        ]
        self.assertEqual(
            [p["display_order"] for p in results],
            list(range(self.first_fixture_order, self.first_fixture_order + 4)),
        )
        self.assertEqual(results[0]["name"], "Tổng Giám đốc kiểm thử")
        self.assertEqual(set(results[0]), {"id", "display_order", "name", "is_public", "note"})

    def test_internal_position_hidden_from_non_admin(self):
        res = self.client_for["pos_viewer"].get(LIST_URL, {"all": "true"})
        names = [p["name"] for p in res.json()]
        self.assertNotIn("Cộng tác viên kiểm thử", names)
        self.assertEqual(len(res.json()), self.baseline_public_count + 3)

    def test_internal_position_visible_to_admin_and_flagged(self):
        for username in ("pos_manager", "pos_root"):
            res = self.client_for[username].get(LIST_URL, {"all": "true"})
            internal = [p for p in res.json() if p["name"] == "Cộng tác viên kiểm thử"]
            self.assertEqual(len(internal), 1)
            self.assertFalse(internal[0]["is_public"])

    def test_search_by_name_case_insensitive(self):
        res = self.client_for["pos_manager"].get(
            LIST_URL, {"search": "trưởng phòng kiểm thử"}
        )
        self.assertEqual([p["name"] for p in res.json()["results"]], ["Trưởng phòng kiểm thử"])

    def test_search_no_match_returns_empty(self):
        res = self.client_for["pos_manager"].get(LIST_URL, {"search": "khong-ton-tai"})
        self.assertEqual(res.json()["count"], 0)

    def test_pagination(self):
        first_order = self.next_order()
        for offset in range(15):
            Position.objects.create(
                display_order=first_order + offset,
                name=f"Chức vụ kiểm thử {offset}",
            )
        page1 = self.client_for["pos_manager"].get(LIST_URL, {"page": 1}).json()
        page2 = self.client_for["pos_manager"].get(LIST_URL, {"page": 2}).json()
        self.assertEqual(page1["count"], Position.objects.count())
        self.assertEqual(len(page1["results"]), 10)
        self.assertEqual(len(page2["results"]), 10)
        self.assertIsNotNone(page1["next"])

    def test_all_true_returns_plain_list(self):
        res = self.client_for["pos_manager"].get(LIST_URL, {"all": "true"})
        self.assertIsInstance(res.json(), list)
        self.assertEqual(len(res.json()), self.baseline_position_count + 4)

    def test_detail_hidden_for_non_admin_404(self):
        internal = Position.objects.get(name="Cộng tác viên kiểm thử")
        self.assertEqual(self.client_for["pos_viewer"].get(detail_url(internal.pk)).status_code, 404)
        self.assertEqual(self.client_for["pos_manager"].get(detail_url(internal.pk)).status_code, 200)

    def test_detail_not_found_404(self):
        res = self.client_for["pos_manager"].get(detail_url(999999))
        self.assertEqual(res.status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(res.json()["detail"], "positions.messages.notFound")


class PositionWriteTests(PositionBaseTestCase):
    def post(self, payload, user="pos_manager"):
        return self.client_for[user].post(LIST_URL, payload, format="json")

    def test_create_without_stt_assigns_max_plus_one(self):
        existing_order = self.next_order(5)
        Position.objects.create(display_order=existing_order, name="Có sẵn")
        res = self.post({"name": "Mới"})
        self.assertEqual(res.status_code, status.HTTP_201_CREATED)
        self.assertEqual(res.json()["display_order"], existing_order + 1)
        self.assertTrue(res.json()["is_public"])

    def test_create_without_stt_uses_next_order_after_seed(self):
        next_order = self.next_order()
        self.assertEqual(self.post({"name": "Đầu tiên"}).json()["display_order"], next_order)

    def test_create_with_explicit_stt_and_fields(self):
        display_order = self.next_order(5)
        res = self.post(
            {"display_order": display_order, "name": "Giám đốc kiểm thử", "is_public": False, "note": "Ghi chú"}
        )
        self.assertEqual(res.status_code, status.HTTP_201_CREATED)
        obj = Position.objects.get(pk=res.json()["id"])
        self.assertEqual(
            (obj.display_order, obj.name, obj.is_public, obj.note),
            (display_order, "Giám đốc kiểm thử", False, "Ghi chú"),
        )

    def test_create_trims_name_and_collapses_spaces(self):
        res = self.post({"name": "   Phó   phòng kiểm thử  "})
        self.assertEqual(res.json()["name"], "Phó phòng kiểm thử")

    def test_create_whitespace_only_name_rejected(self):
        res = self.post({"name": "     "})
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(res.json()["name"], ["positions.messages.nameRequired"])

    def test_create_missing_name_rejected(self):
        res = self.post({"display_order": self.next_order()})
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("name", res.json())

    def test_create_name_too_long_rejected(self):
        res = self.post({"name": "x" * 256})
        self.assertEqual(res.json()["name"], ["positions.messages.nameTooLong"])

    def test_create_duplicate_name_after_trim_and_case_rejected(self):
        Position.objects.create(display_order=self.next_order(), name="Trưởng phòng kiểm thử")
        for variant in (
            "trưởng phòng kiểm thử", "  TRƯỞNG PHÒNG KIỂM THỬ  ", "Trưởng   phòng kiểm thử"
        ):
            res = self.post({"name": variant})
            self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST, variant)
            self.assertEqual(res.json()["name"], ["positions.messages.nameExists"])

    def test_create_duplicate_stt_rejected(self):
        duplicate_order = self.next_order()
        Position.objects.create(display_order=duplicate_order, name="A")
        res = self.post({"display_order": duplicate_order, "name": "B"})
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(res.json()["display_order"], ["positions.messages.sttExists"])

    def test_create_invalid_stt_rejected(self):
        for bad in (0, -3, "abc", 1.5):
            res = self.post({"display_order": bad, "name": f"X{bad}"})
            self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST, bad)
            self.assertEqual(res.json()["display_order"], ["positions.messages.sttInvalid"])

    def test_create_requires_add_permission(self):
        self.assertEqual(self.post({"name": "A"}, user="pos_viewer").status_code, 403)
        self.assertEqual(self.post({"name": "A"}, user="pos_noperm").status_code, 403)
        self.assertEqual(self.others().count(), 0)

    def test_create_unauthenticated_401(self):
        self.assertEqual(self.anon.post(LIST_URL, {"name": "A"}, format="json").status_code, 401)

    def test_superuser_can_write_without_explicit_permissions(self):
        self.assertEqual(self.post({"name": "Root tạo"}, user="pos_root").status_code, 201)

    def test_auto_stt_retries_when_number_taken_concurrently(self):
        """Hai request cùng chọn MAX+1: lần lưu đầu va unique, phải thử lại với số kế tiếp."""
        used_order = self.next_order()
        Position.objects.create(display_order=used_order, name="Có sẵn")
        calls = {"n": 0}
        real_next = Position.get_next_display_order

        def flaky_next():
            calls["n"] += 1
            if calls["n"] == 1:
                return used_order  # số đã bị người khác lấy
            return real_next()

        with mock.patch.object(Position, "get_next_display_order", side_effect=flaky_next):
            res = self.post({"name": "Sau retry"})
        self.assertEqual(res.status_code, status.HTTP_201_CREATED)
        self.assertEqual(res.json()["display_order"], used_order + 1)
        self.assertEqual(calls["n"], 2)

    def test_put_updates_all_fields(self):
        starting_order = self.next_order()
        target_order = starting_order + 1
        pos = Position.objects.create(display_order=starting_order, name="Cũ")
        res = self.client_for["pos_manager"].put(
            detail_url(pos.pk),
            {"display_order": target_order, "name": "Mới", "is_public": False, "note": "x"},
            format="json",
        )
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        pos.refresh_from_db()
        self.assertEqual((pos.display_order, pos.name, pos.is_public, pos.note), (target_order, "Mới", False, "x"))

    def test_put_without_stt_keeps_current_stt(self):
        starting_order = self.next_order()
        pos = Position.objects.create(display_order=starting_order, name="Giữ STT")
        res = self.client_for["pos_manager"].put(detail_url(pos.pk), {"name": "Đổi tên"}, format="json")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        pos.refresh_from_db()
        self.assertEqual(pos.display_order, starting_order)

    def test_put_same_name_and_stt_on_itself_allowed(self):
        starting_order = self.next_order()
        pos = Position.objects.create(display_order=starting_order, name="Chính nó")
        res = self.client_for["pos_manager"].put(
            detail_url(pos.pk), {"display_order": starting_order, "name": "chính nó"}, format="json"
        )
        self.assertEqual(res.status_code, status.HTTP_200_OK)

    def test_put_duplicate_name_or_stt_of_other_rejected(self):
        first_order = self.next_order()
        Position.objects.create(display_order=first_order, name="Khác")
        pos = Position.objects.create(display_order=first_order + 1, name="Tôi")
        client = self.client_for["pos_manager"]
        res = client.put(detail_url(pos.pk), {"name": "khác"}, format="json")
        self.assertEqual(res.json()["name"], ["positions.messages.nameExists"])
        res = client.put(detail_url(pos.pk), {"display_order": first_order, "name": "Tôi"}, format="json")
        self.assertEqual(res.json()["display_order"], ["positions.messages.sttExists"])

    def test_patch_partial(self):
        pos = Position.objects.create(display_order=self.next_order(), name="Giữ tên", note="n")
        res = self.client_for["pos_manager"].patch(detail_url(pos.pk), {"is_public": False}, format="json")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        pos.refresh_from_db()
        self.assertEqual((pos.name, pos.note, pos.is_public), ("Giữ tên", "n", False))

    def test_update_requires_change_permission(self):
        pos = Position.objects.create(display_order=self.next_order(), name="X")
        res = self.client_for["pos_viewer"].put(detail_url(pos.pk), {"name": "Y"}, format="json")
        self.assertEqual(res.status_code, 403)
        res = self.client_for["pos_viewer"].patch(detail_url(pos.pk), {"name": "Y"}, format="json")
        self.assertEqual(res.status_code, 403)
        pos.refresh_from_db()
        self.assertEqual(pos.name, "X")

    def test_update_not_found(self):
        res = self.client_for["pos_manager"].put(detail_url(999999), {"name": "Y"}, format="json")
        self.assertEqual(res.status_code, 404)

    def test_delete_success(self):
        pos = Position.objects.create(display_order=self.next_order(), name="Xóa tôi")
        res = self.client_for["pos_manager"].delete(detail_url(pos.pk))
        self.assertEqual(res.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(Position.objects.filter(pk=pos.pk).exists())

    def test_delete_requires_permission(self):
        pos = Position.objects.create(display_order=self.next_order(), name="Giữ")
        self.assertEqual(self.client_for["pos_viewer"].delete(detail_url(pos.pk)).status_code, 403)
        self.assertEqual(self.anon.delete(detail_url(pos.pk)).status_code, 401)
        self.assertTrue(Position.objects.filter(pk=pos.pk).exists())

    def test_delete_position_in_use_by_employee_400(self):
        khoi = KhoiKhuVuc.objects.create(ten="Khối T")
        latest_stt = DonVi.objects.order_by("-stt_uu_tien").values_list("stt_uu_tien", flat=True).first() or 0
        don_vi = DonVi.objects.create(ten="Đơn vị T", khoi=khoi, stt_uu_tien=latest_stt + 1)
        pos = Position.objects.create(display_order=self.next_order(), name="Đang dùng")
        Employee.objects.create_user(
            username="pos_emp", password="Test@12345", ho_ten="NV", chuc_vu=pos, don_vi=don_vi
        )
        res = self.client_for["pos_manager"].delete(detail_url(pos.pk))
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(res.json()["detail"], "positions.messages.inUse")
        self.assertTrue(Position.objects.filter(pk=pos.pk).exists())


class PositionTemplateTests(PositionBaseTestCase):
    def test_template_download(self):
        res = self.client_for["pos_manager"].get(TEMPLATE_URL)
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertIn("spreadsheetml.sheet", res["Content-Type"])
        self.assertIn("mau_import_chuc_vu.xlsx", res["Content-Disposition"])
        wb = openpyxl.load_workbook(BytesIO(res.content))
        self.assertEqual([c.value for c in wb.worksheets[0][1]], HEADERS)
        self.assertEqual(wb.worksheets[0].max_row, 1)

    def test_template_requires_import_permission(self):
        self.assertEqual(self.client_for["pos_viewer"].get(TEMPLATE_URL).status_code, 403)
        self.assertEqual(self.anon.get(TEMPLATE_URL).status_code, 401)

    def test_downloaded_template_is_importable_after_filling(self):
        res = self.client_for["pos_manager"].get(TEMPLATE_URL)
        wb = openpyxl.load_workbook(BytesIO(res.content))
        wb.worksheets[0].append([None, "Điền từ mẫu", "Có", ""])
        buffer = BytesIO()
        wb.save(buffer)
        upload = SimpleUploadedFile("filled.xlsx", buffer.getvalue())
        res = self.client_for["pos_manager"].post(IMPORT_URL, {"file": upload}, format="multipart")
        self.assertEqual(res.json()["imported"], 1)


class PositionImportTests(PositionBaseTestCase):
    def do_import(self, upload, user="pos_manager"):
        return self.client_for[user].post(IMPORT_URL, {"file": upload}, format="multipart")

    def test_import_valid_rows(self):
        first_order = self.next_order()
        res = self.do_import(
            make_xlsx(
                [
                    [first_order, "Tổng Giám đốc import kiểm thử", "Có", "Cấp cao nhất"],
                    [first_order + 1, "Trưởng phòng import kiểm thử", "Không", ""],
                    [None, "Chuyên viên import kiểm thử", None, None],
                ]
            )
        )
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.json(), {"total": 3, "imported": 3, "failed": 0, "errors": []})
        self.assertEqual(Position.objects.get(name="Trưởng phòng import kiểm thử").is_public, False)
        self.assertEqual(Position.objects.get(name="Tổng Giám đốc import kiểm thử").note, "Cấp cao nhất")
        self.assertEqual(Position.objects.get(name="Chuyên viên import kiểm thử").display_order, first_order + 2)

    def test_import_blank_stt_avoids_explicit_stt_later_in_file(self):
        first_order = self.next_order()
        res = self.do_import(
            make_xlsx([[None, "Tự động", "Có", ""], [first_order, "Cố định", "Có", ""]])
        )
        self.assertEqual(res.json()["imported"], 2)
        self.assertEqual(Position.objects.get(name="Cố định").display_order, first_order)
        self.assertEqual(Position.objects.get(name="Tự động").display_order, first_order + 1)

    def test_import_reports_specific_error_per_row(self):
        first_order = self.next_order()
        Position.objects.create(display_order=first_order, name="Đã có")
        Position.objects.create(display_order=first_order + 1, name="Trùng STT")
        res = self.do_import(
            make_xlsx(
                [
                    [None, "Hợp lệ", "Có", ""],  # row 2
                    [None, "", "Có", ""],  # row 3: thiếu tên
                    [None, "  đã   có  ", "Có", ""],  # row 4: trùng DB sau trim
                    [None, "Trùng trong file", "Có", ""],  # row 5 ok
                    [None, "trùng trong FILE", "Có", ""],  # row 6: trùng trong file
                    [first_order, "STT trùng DB", "Có", ""],  # row 7
                    ["abc", "STT sai", "Có", ""],  # row 8
                    [0, "STT không dương", "Có", ""],  # row 9
                    [None, "Công khai sai", "Có thể", ""],  # row 10
                    [50, "A", "Có", ""],  # row 11 ok
                    [50, "B", "Có", ""],  # row 12: trùng STT trong file
                ]
            )
        )
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        body = res.json()
        self.assertEqual(body["total"], 11)
        self.assertEqual(body["imported"], 3)
        self.assertEqual(body["failed"], 8)
        errors = {(e["row"], e["field"]): e["code"] for e in body["errors"]}
        self.assertEqual(
            errors,
            {
                (3, "name"): "positions.messages.nameRequired",
                (4, "name"): "positions.messages.nameExists",
                (6, "name"): "positions.import.errors.duplicateNameInFile",
                (7, "display_order"): "positions.messages.sttExists",
                (8, "display_order"): "positions.import.errors.invalidStt",
                (9, "display_order"): "positions.import.errors.invalidStt",
                (10, "is_public"): "positions.import.errors.invalidPublic",
                (12, "display_order"): "positions.import.errors.duplicateSttInFile",
            },
        )
        self.assertEqual([e["row"] for e in body["errors"]], sorted(e["row"] for e in body["errors"]))
        self.assertTrue(Position.objects.filter(name="Hợp lệ").exists())
        self.assertFalse(Position.objects.filter(name="STT sai").exists())

    def test_import_row_with_two_errors_counts_as_one_failed_row(self):
        duplicate_order = self.next_order()
        Position.objects.create(display_order=duplicate_order, name="Đã có")
        res = self.do_import(make_xlsx([[duplicate_order, "đã có", "Có", ""]]))
        body = res.json()
        self.assertEqual(body["failed"], 1)
        self.assertEqual(len(body["errors"]), 2)
        self.assertEqual(body["imported"], 0)

    def test_import_accepts_float_and_text_stt(self):
        first_order = self.next_order()
        res = self.do_import(
            make_xlsx([
                [float(first_order), "Số thực kiểm thử", "Có", ""],
                [str(first_order + 1), "Số chữ kiểm thử", "Có", ""],
            ])
        )
        self.assertEqual(res.json()["imported"], 2)
        self.assertEqual(
            Position.objects.get(name="Số chữ kiểm thử").display_order, first_order + 1
        )

    def test_import_skips_blank_rows_but_keeps_row_numbers(self):
        res = self.do_import(make_xlsx([[None, None, None, None], [None, "", "Có", ""], [None, "Sau dòng trống", "Có", ""]]))
        body = res.json()
        self.assertEqual(body["total"], 2)
        self.assertEqual([e["row"] for e in body["errors"]], [3])

    def test_import_header_tolerates_case_and_required_marker(self):
        res = self.do_import(
            make_xlsx([[None, "Header khác", "Có", ""]], header=["stt ưu tiên", "TÊN CHỨC VỤ", "hiển thị công khai (có/không)", "ghi chú"])
        )
        self.assertEqual(res.json()["imported"], 1)

    def test_import_wrong_template_400(self):
        res = self.do_import(make_xlsx([["a", "b", "c", "d"]], header=["Cột 1", "Cột 2", "Cột 3", "Cột 4"]))
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(res.json()["detail"], "positions.import.errors.invalidTemplate")

    def test_import_header_only_is_empty_400(self):
        res = self.do_import(make_xlsx([]))
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(res.json()["detail"], "positions.import.errors.emptyFile")

    def test_import_not_an_excel_file_400(self):
        res = self.do_import(SimpleUploadedFile("fake.xlsx", b"not an excel file at all"))
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(res.json()["detail"], "positions.import.errors.invalidFile")

    def test_import_wrong_extension_400(self):
        res = self.do_import(SimpleUploadedFile("data.csv", b"a,b,c"))
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(res.json()["detail"], "positions.import.errors.invalidFile")

    def test_import_missing_file_400(self):
        res = self.client_for["pos_manager"].post(IMPORT_URL, {}, format="multipart")
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(res.json()["detail"], "positions.import.errors.fileRequired")

    def test_import_too_many_rows_400(self):
        rows = [[None, f"Chức vụ {i}", "Có", ""] for i in range(MAX_ROWS + 1)]
        res = self.do_import(make_xlsx(rows))
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(res.json()["detail"], "positions.import.errors.tooManyRows")
        self.assertEqual(self.others().count(), 0)

    def test_import_requires_import_permission(self):
        upload = make_xlsx([[None, "Không được", "Có", ""]])
        self.assertEqual(self.do_import(upload, user="pos_viewer").status_code, 403)
        self.assertEqual(self.anon.post(IMPORT_URL, {"file": make_xlsx([])}, format="multipart").status_code, 401)
        self.assertEqual(self.others().count(), 0)

    def test_import_note_is_trimmed_and_name_collapsed(self):
        res = self.do_import(
            make_xlsx([[None, "  Phó   Giám  đốc kiểm thử ", "Có", "  ghi chú  "]])
        )
        self.assertEqual(res.json()["imported"], 1)
        obj = Position.objects.get(name="Phó Giám đốc kiểm thử")
        self.assertEqual(obj.note, "ghi chú")
