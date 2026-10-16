import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from django.db import IntegrityError, transaction
from django.db.models import Q, ProtectedError
from django.http import HttpResponse
from rest_framework import status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.pagination import PageNumberPagination
from rest_framework.permissions import BasePermission, IsAuthenticated
from rest_framework.response import Response

from .excel_i18n import get_excel_translations
from .models import KhoiKhuVuc, DonVi
from .references import unit_in_use
from .serializers import KhoiKhuVucSerializer, DonViSerializer


def _invalid_region_filter(region_id):
    """?region_id= phải là số nguyên dương trong khoảng BIGINT; sai thì trả 400 thay vì lỗi 500."""
    if region_id and (not region_id.isdigit() or int(region_id) > 9223372036854775807):
        return Response(
            {"detail": "organizations.messages.regionNotFound", "key": "organizations.messages.regionNotFound"},
            status=status.HTTP_400_BAD_REQUEST,
        )
    return None


class StandardPagination(PageNumberPagination):
    page_size = 10
    page_size_query_param = "page_size"
    max_page_size = 100


# Khối/Khu vực và Đơn vị dùng chung bộ quyền *_donvi của DDL (DDL không có quyền riêng cho Khối).
LIST_REQUIRED_PERM = {"GET": "view_donvi", "POST": "add_donvi"}
DETAIL_REQUIRED_PERM = {
    "GET": "view_donvi",
    "PUT": "change_donvi",
    "PATCH": "change_donvi",
    "DELETE": "delete_donvi",
}


def _permission_denied():
    key = "organizations.messages.permissionDenied"
    return Response({"detail": key, "key": key}, status=status.HTTP_403_FORBIDDEN)


class CanExportUnits(BasePermission):
    message = "organizations.messages.exportPermissionDenied"

    def has_permission(self, request, view):
        return bool(
            request.user
            and request.user.is_authenticated
            and request.user.has_perm("organizations.export_donvi")
        )


@api_view(["GET", "POST"])
@permission_classes([IsAuthenticated])
def region_list_create_view(request):
    if not request.user.has_perm(LIST_REQUIRED_PERM[request.method]):
        return _permission_denied()

    if request.method == "GET":
        queryset = KhoiKhuVuc.objects.all().order_by("ten")
        search_query = request.query_params.get("search", "").strip()
        if search_query:
            queryset = queryset.filter(ten__icontains=search_query)

        # Trả về toàn bộ danh sách không phân trang nếu có tham số all=true (dành cho dropdown/select)
        if request.query_params.get("all", "").lower() == "true":
            serializer = KhoiKhuVucSerializer(queryset, many=True)
            return Response(serializer.data, status=status.HTTP_200_OK)

        # Phân trang chuẩn server-side
        paginator = StandardPagination()
        page = paginator.paginate_queryset(queryset, request)
        if page is not None:
            serializer = KhoiKhuVucSerializer(page, many=True)
            return paginator.get_paginated_response(serializer.data)

        serializer = KhoiKhuVucSerializer(queryset, many=True)
        return Response(serializer.data, status=status.HTTP_200_OK)

    elif request.method == "POST":
        serializer = KhoiKhuVucSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


@api_view(["GET", "PUT", "PATCH", "DELETE"])
@permission_classes([IsAuthenticated])
def region_detail_view(request, pk):
    if not request.user.has_perm(DETAIL_REQUIRED_PERM[request.method]):
        return _permission_denied()

    try:
        region = KhoiKhuVuc.objects.get(pk=pk)
    except KhoiKhuVuc.DoesNotExist:
        return Response(
            {"detail": "organizations.messages.regionNotFound", "key": "organizations.messages.regionNotFound"},
            status=status.HTTP_404_NOT_FOUND,
        )

    if request.method == "GET":
        serializer = KhoiKhuVucSerializer(region)
        return Response(serializer.data, status=status.HTTP_200_OK)

    elif request.method in ["PUT", "PATCH"]:
        partial = (request.method == "PATCH")
        serializer = KhoiKhuVucSerializer(region, data=request.data, partial=partial)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data, status=status.HTTP_200_OK)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    elif request.method == "DELETE":
        if region.don_vi_list.exists():
            return Response(
                {"detail": "organizations.messages.regionHasUnits", "key": "organizations.messages.regionHasUnits"},
                status=status.HTTP_400_BAD_REQUEST,
            )
        try:
            region.delete()
            return Response(
                {"detail": "organizations.messages.deleteSuccess", "key": "organizations.messages.deleteSuccess"},
                status=status.HTTP_204_NO_CONTENT,
            )
        except ProtectedError:
            return Response(
                {"detail": "organizations.messages.protectedError", "key": "organizations.messages.protectedError"},
                status=status.HTTP_400_BAD_REQUEST,
            )


@api_view(["GET", "POST"])
@permission_classes([IsAuthenticated])
def unit_list_create_view(request):
    if not request.user.has_perm(LIST_REQUIRED_PERM[request.method]):
        return _permission_denied()

    if request.method == "GET":
        queryset = DonVi.objects.select_related("khoi").all().order_by("stt_uu_tien", "ten")

        region_id = request.query_params.get("region_id")
        invalid = _invalid_region_filter(region_id)
        if invalid:
            return invalid
        if region_id:
            queryset = queryset.filter(khoi_id=region_id)

        search_query = request.query_params.get("search", "").strip()
        if search_query:
            queryset = queryset.filter(
                Q(ten__icontains=search_query) |
                Q(dien_thoai__icontains=search_query) |
                Q(email__icontains=search_query) |
                Q(dia_chi__icontains=search_query)
            )

        paginator = StandardPagination()
        page = paginator.paginate_queryset(queryset, request)
        if page is not None:
            serializer = DonViSerializer(page, many=True)
            return paginator.get_paginated_response(serializer.data)

        serializer = DonViSerializer(queryset, many=True)
        return Response(serializer.data, status=status.HTTP_200_OK)

    elif request.method == "POST":
        serializer = DonViSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


@api_view(["GET", "PUT", "PATCH", "DELETE"])
@permission_classes([IsAuthenticated])
def unit_detail_view(request, pk):
    if not request.user.has_perm(DETAIL_REQUIRED_PERM[request.method]):
        return _permission_denied()

    try:
        unit = DonVi.objects.select_related("khoi").get(pk=pk)
    except DonVi.DoesNotExist:
        return Response(
            {"detail": "organizations.messages.unitNotFound", "key": "organizations.messages.unitNotFound"},
            status=status.HTTP_404_NOT_FOUND,
        )

    if request.method == "GET":
        serializer = DonViSerializer(unit)
        return Response(serializer.data, status=status.HTTP_200_OK)

    elif request.method in ["PUT", "PATCH"]:
        partial = (request.method == "PATCH")
        serializer = DonViSerializer(unit, data=request.data, partial=partial)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data, status=status.HTTP_200_OK)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    elif request.method == "DELETE":
        in_use = Response(
            {"detail": "organizations.messages.unitInUse", "key": "organizations.messages.unitInUse"},
            status=status.HTTP_400_BAD_REQUEST,
        )
        if unit_in_use(unit.pk):
            return in_use
        try:
            with transaction.atomic():
                unit.delete()
        except ProtectedError:
            return Response(
                {"detail": "organizations.messages.unitHasEmployees", "key": "organizations.messages.unitHasEmployees"},
                status=status.HTTP_400_BAD_REQUEST,
            )
        except IntegrityError:  # DB dựng từ DDL chặn bằng ON DELETE RESTRICT
            return in_use
        return Response(
            {"detail": "organizations.messages.deleteSuccess", "key": "organizations.messages.deleteSuccess"},
            status=status.HTTP_204_NO_CONTENT,
        )


@api_view(["GET"])
@permission_classes([IsAuthenticated, CanExportUnits])
def unit_export_excel_view(request):
    lang = request.query_params.get("lang", "vi").lower()
    i18n = get_excel_translations(lang)

    queryset = DonVi.objects.select_related("khoi").all().order_by("stt_uu_tien", "ten")

    region_id = request.query_params.get("region_id")
    invalid = _invalid_region_filter(region_id)
    if invalid:
        return invalid
    if region_id:
        queryset = queryset.filter(khoi_id=region_id)

    search_query = request.query_params.get("search", "").strip()
    if search_query:
        queryset = queryset.filter(
            Q(ten__icontains=search_query) |
            Q(dien_thoai__icontains=search_query) |
            Q(email__icontains=search_query) |
            Q(dia_chi__icontains=search_query)
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = i18n["sheet_title"]

    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    headers = i18n["headers"]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    for row_idx, unit in enumerate(queryset, start=2):
        status_text = i18n["public"] if unit.is_public else i18n["internal"]
        row_data = [
            unit.stt_uu_tien,
            unit.ten,
            unit.khoi.ten if unit.khoi else "",
            unit.dien_thoai,
            unit.email,
            unit.dia_chi,
            status_text,
            unit.ghi_chu,
        ]
        ws.append(row_data)

        for col_idx in range(1, len(row_data) + 1):
            c = ws.cell(row=row_idx, column=col_idx)
            c.border = thin_border
            if col_idx in [1, 4, 7]:
                c.alignment = center_align
            else:
                c.alignment = left_align

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = i18n["filename"]
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
