import io
import openpyxl
from unittest.mock import patch
from django.core.exceptions import ValidationError
from django.db import IntegrityError
from django.test import TestCase
from rest_framework.test import APIClient
from rest_framework import status
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission

from organizations.models import KhoiKhuVuc, DonVi


class OrganizationModelAndAPITest(TestCase):
    """
    Test suite kiểm tra phân hệ Organizations: Model, Validation Keys, và API Endpoints.
    """

    def setUp(self):
        self.region = KhoiKhuVuc.objects.create(ten="Khối Cơ quan")
        self.unit = DonVi.objects.create(
            khoi=self.region,
            stt_uu_tien=1,
            ten="Phòng Kế toán",
            dien_thoai="02838123456",
            email="ketoan@itc.edu.vn",
        )
        self.user = get_user_model().objects.create_user(
            username="testuser",
            password="testpassword123",
            email="test@itc.edu.vn",
            don_vi=self.unit,
        )
        self.export_permission = Permission.objects.get(codename="export_donvi")
        self.user.user_permissions.add(
            *Permission.objects.filter(
                codename__in=["view_donvi", "add_donvi", "change_donvi", "delete_donvi", "export_donvi"]
            )
        )
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

    def test_khoi_khuvuc_creation(self):
        self.assertEqual(str(self.region), "Khối Cơ quan")
        self.assertEqual(self.region.don_vi_list.count(), 1)

    def test_donvi_creation(self):
        self.assertEqual(self.unit.stt_uu_tien, 1)
        self.assertTrue(self.unit.is_public)

    def test_auto_priority_assignment(self):
        data = {
            "khoi": self.region.id,
            "ten": "Phòng Tổ chức Cán bộ",
            "dien_thoai": "0912345678",
            "email": "tccb@itc.edu.vn",
        }
        response = self.client.post("/api/organizations/units/", data, format="json")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["stt_uu_tien"], 2)

    def test_validation_key_on_blank_name(self):
        data = {
            "khoi": self.region.id,
            "ten": "   ",
        }
        response = self.client.post("/api/organizations/units/", data, format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("organizations.messages.nameRequired", str(response.data))

    def test_model_save_rejects_invalid_values(self):
        """Kiểm tra model tự chặn dữ liệu sai khi lưu trực tiếp, không qua API."""
        with self.assertRaises(ValidationError) as region_context:
            KhoiKhuVuc(ten="   ").save()
        self.assertIn("organizations.messages.nameRequired", str(region_context.exception))

        with self.assertRaises(ValidationError) as unit_context:
            DonVi(
                khoi=self.region,
                stt_uu_tien=2,
                ten="Đơn vị không hợp lệ",
                dien_thoai="12345678",
            ).save()
        self.assertIn("organizations.messages.phoneInvalid", str(unit_context.exception))

    def test_prevent_delete_region_with_units(self):
        response = self.client.delete(f"/api/organizations/regions/{self.region.id}/")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(response.data.get("key"), "organizations.messages.regionHasUnits")

    def test_export_excel_endpoint_and_content(self):
        """Kiểm tra tệp Excel được sinh ra với tiêu đề song ngữ và nội dung chính xác."""
        # 1. Kiểm tra xuất tiếng Việt
        res_vi = self.client.get("/api/organizations/units/export/?lang=vi")
        self.assertEqual(res_vi.status_code, status.HTTP_200_OK)
        self.assertEqual(
            res_vi["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        wb_vi = openpyxl.load_workbook(io.BytesIO(res_vi.content))
        ws_vi = wb_vi.active
        self.assertEqual(ws_vi.title, "Danh Sách Đơn Vị")
        headers_vi = [cell.value for cell in ws_vi[1]]
        self.assertIn("STT Ưu Tiên", headers_vi)
        self.assertIn("Tên Đơn Vị / Phòng Ban", headers_vi)
        # Kiểm tra nội dung hàng 2
        row2_vi = [cell.value for cell in ws_vi[2]]
        self.assertEqual(row2_vi[1], "Phòng Kế toán")

        # 2. Kiểm tra xuất tiếng Anh
        res_en = self.client.get("/api/organizations/units/export/?lang=en")
        self.assertEqual(res_en.status_code, status.HTTP_200_OK)
        wb_en = openpyxl.load_workbook(io.BytesIO(res_en.content))
        ws_en = wb_en.active
        self.assertEqual(ws_en.title, "Departments Directory")
        headers_en = [cell.value for cell in ws_en[1]]
        self.assertIn("Priority Index", headers_en)
        self.assertIn("Unit / Department Name", headers_en)

    def test_auto_priority_on_blank_string(self):
        """Kiểm tra khi frontend gửi chuỗi rỗng stt_uu_tien="", hệ thống vẫn tự động gán max + 1."""
        data = {
            "khoi": self.region.id,
            "stt_uu_tien": "",
            "ten": "Đơn Vị STT Chuỗi Rỗng",
        }
        response = self.client.post("/api/organizations/units/", data, format="json")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["stt_uu_tien"], 2)

    def test_auto_priority_retries_after_integrity_error(self):
        """Mô phỏng xung đột đồng thời và kiểm tra hệ thống tự thử lại việc cấp STT."""
        data = {
            "khoi": self.region.id,
            "stt_uu_tien": "",
            "ten": "Đơn Vị Retry STT",
        }
        original_create = DonVi.objects.create
        attempts = {"count": 0}

        def create_with_conflict(**kwargs):
            attempts["count"] += 1
            if attempts["count"] == 1:
                raise IntegrityError("duplicate priority")
            return original_create(**kwargs)

        with patch.object(
            DonVi.objects,
            "create",
            side_effect=create_with_conflict,
        ):
            response = self.client.post("/api/organizations/units/", data, format="json")

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["stt_uu_tien"], 2)
        self.assertEqual(attempts["count"], 2)

    def test_export_requires_permission(self):
        """Kiểm tra người dùng đã đăng nhập nhưng thiếu quyền export bị từ chối."""
        self.user.user_permissions.remove(self.export_permission)

        response = self.client.get("/api/organizations/units/export/")

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(response.data["detail"], "organizations.messages.exportPermissionDenied")

    def test_duplicate_priority_rejection(self):
        """Kiểm tra không cho phép trùng STT ưu tiên đã tồn tại."""
        data = {
            "khoi": self.region.id,
            "stt_uu_tien": 1,
            "ten": "Đơn vị Trùng STT",
        }
        response = self.client.post("/api/organizations/units/", data, format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("organizations.messages.priorityExists", str(response.data))

    def test_invalid_phone_number_rejection(self):
        """Kiểm tra regex số điện thoại Việt Nam hợp lệ."""
        data = {
            "khoi": self.region.id,
            "ten": "Đơn vị SĐT Lỗi",
            "dien_thoai": "12345678",  # Không bắt đầu bằng 0
        }
        response = self.client.post("/api/organizations/units/", data, format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("organizations.messages.phoneInvalid", str(response.data))

    def test_invalid_email_format(self):
        """Kiểm tra định dạng email không hợp lệ."""
        data = {
            "khoi": self.region.id,
            "ten": "Đơn vị Email Lỗi",
            "email": "not-a-valid-email",
        }
        response = self.client.post("/api/organizations/units/", data, format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("organizations.messages.emailInvalid", str(response.data))

    def test_region_pagination_and_all_param(self):
        """Kiểm tra phân trang mặc định nhiều trang và tham số ?all=true cho Khối/Khu vực."""
        # Tạo thêm 12 Khối để có tổng cộng 13 Khối (vượt page_size = 10)
        KhoiKhuVuc.objects.bulk_create([
            KhoiKhuVuc(ten=f"Khối Thử Nghiệm Số {i:02d}") for i in range(2, 14)
        ])

        # 1. Gọi trang 1 -> phải có 10 bản ghi và có liên kết trang tiếp theo (next)
        res_page1 = self.client.get("/api/organizations/regions/?page=1")
        self.assertEqual(res_page1.status_code, status.HTTP_200_OK)
        self.assertEqual(res_page1.data["count"], 13)
        self.assertEqual(len(res_page1.data["results"]), 10)
        self.assertIsNotNone(res_page1.data["next"])

        # 2. Gọi trang 2 -> phải có 3 bản ghi còn lại và có liên kết trang trước (previous)
        res_page2 = self.client.get("/api/organizations/regions/?page=2")
        self.assertEqual(res_page2.status_code, status.HTTP_200_OK)
        self.assertEqual(len(res_page2.data["results"]), 3)
        self.assertIsNotNone(res_page2.data["previous"])

        # 3. Gọi với ?all=true -> phải trả về toàn bộ 13 bản ghi dạng mảng không phân trang
        res_all = self.client.get("/api/organizations/regions/?all=true")
        self.assertEqual(res_all.status_code, status.HTTP_200_OK)
        self.assertIsInstance(res_all.data, list)
        self.assertEqual(len(res_all.data), 13)

    def test_unit_search_by_address(self):
        """Kiểm tra tìm kiếm đơn vị theo địa chỉ."""
        self.unit.dia_chi = "254 Nguyễn Văn Linh, Đà Nẵng"
        self.unit.save()

        response = self.client.get("/api/organizations/units/?search=Nguyễn+Văn+Linh")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["count"], 1)

        response_not_found = self.client.get("/api/organizations/units/?search=Hà+Nội")
        self.assertEqual(response_not_found.status_code, status.HTTP_200_OK)
        self.assertEqual(response_not_found.data["count"], 0)

    def test_export_search_by_address(self):
        """Kiểm tra xuất file Excel lọc đúng địa chỉ và loại trừ các bản ghi không khớp."""
        # Đơn vị 1 khớp từ khóa Đà Nẵng
        self.unit.dia_chi = "254 Nguyễn Văn Linh, Đà Nẵng"
        self.unit.save()

        # Tạo Đơn vị 2 có địa chỉ ở Hà Nội (không khớp từ khóa tìm kiếm)
        DonVi.objects.create(
            khoi=self.region,
            stt_uu_tien=99,
            ten="Văn Phòng Đại Diện",
            dia_chi="123 Phố Cầu Giấy, Hà Nội",
        )

        # 1. Xuất file với từ khóa 'Nguyễn Văn Linh'
        response = self.client.get("/api/organizations/units/export/?search=Nguyễn+Văn+Linh")
        self.assertEqual(response.status_code, status.HTTP_200_OK)

        # Mở file Excel từ luồng byte và đọc các dòng
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))

        # Phải có đúng 2 dòng: dòng 1 là tiêu đề cột, dòng 2 là đơn vị khớp địa chỉ
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][1], "Phòng Kế toán")
        self.assertIn("Nguyễn Văn Linh", rows[1][5])

        # 2. Xuất file với từ khóa không tồn tại -> chỉ có đúng 1 dòng tiêu đề cột
        res_empty = self.client.get("/api/organizations/units/export/?search=Địa+Chỉ+Không+Tồn+Tại")
        wb_empty = openpyxl.load_workbook(io.BytesIO(res_empty.content))
        ws_empty = wb_empty.active
        rows_empty = list(ws_empty.iter_rows(values_only=True))
        self.assertEqual(len(rows_empty), 1)

    def test_unit_update_and_delete(self):
        """Kiểm tra cập nhật và xóa đơn vị thành công."""
        # Cập nhật PUT
        update_data = {
            "khoi": self.region.id,
            "stt_uu_tien": 1,
            "ten": "Phòng Kế toán Đổi Tên",
            "dien_thoai": "02838999999",
            "email": "ketoan_new@itc.edu.vn",
            "is_public": False,
        }
        res_put = self.client.put(f"/api/organizations/units/{self.unit.id}/", update_data, format="json")
        self.assertEqual(res_put.status_code, status.HTTP_200_OK)
        self.assertEqual(res_put.data["ten"], "Phòng Kế toán Đổi Tên")
        self.assertFalse(res_put.data["is_public"])

        # Xóa DELETE (đơn vị không có nhân viên)
        spare = DonVi.objects.create(khoi=self.region, stt_uu_tien=2, ten="Phòng Trống")
        res_delete = self.client.delete(f"/api/organizations/units/{spare.id}/")
        self.assertEqual(res_delete.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(DonVi.objects.filter(pk=spare.id).exists())

    def test_prevent_delete_unit_with_employees(self):
        """Đơn vị còn nhân viên trực thuộc không được xóa (400, không phải 500)."""
        res = self.client.delete(f"/api/organizations/units/{self.unit.id}/")
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(res.data["detail"], "organizations.messages.unitHasEmployees")
        self.assertTrue(DonVi.objects.filter(pk=self.unit.id).exists())

    def test_invalid_input_returns_400_not_500(self):
        """Đầu vào vượt giới hạn cột DDL hoặc sai kiểu phải trả 400 kèm key, không lỗi 500."""
        cases = [
            ("get", "/api/organizations/units/?region_id=abc", None, "detail", "organizations.messages.regionNotFound"),
            ("get", "/api/organizations/units/export/?region_id=1x", None, "detail", "organizations.messages.regionNotFound"),
            ("post", "/api/organizations/regions/", {"ten": "K" * 256}, "ten", "organizations.messages.nameTooLong"),
            ("post", "/api/organizations/regions/", {"ten": "Khối lớn", "thu_tu_hien_thi": 2147483648}, "thu_tu_hien_thi",
             "organizations.messages.displayOrderInvalid"),
            ("post", "/api/organizations/units/", {"khoi": self.region.id, "ten": "D" * 256}, "ten",
             "organizations.messages.nameTooLong"),
            ("post", "/api/organizations/units/", {"khoi": self.region.id, "ten": "Đơn vị lớn", "stt_uu_tien": 2147483648},
             "stt_uu_tien", "organizations.messages.priorityTooLarge"),
            ("post", "/api/organizations/units/", {"khoi": self.region.id, "ten": "Đơn vị địa chỉ", "dia_chi": "a" * 256},
             "dia_chi", "organizations.messages.addressTooLong"),
        ]
        for method, url, data, key, message in cases:
            with self.subTest(url=url, key=key):
                res = getattr(self.client, method)(url, data, format="json") if data else getattr(self.client, method)(url)
                self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
                value = res.data[key]
                self.assertEqual(value if isinstance(value, str) else value[0], message)

    def test_unit_nullable_columns_accept_null(self):
        """DDL cho dia_chi, dien_thoai, email, ghi_chu nhận NULL."""
        res = self.client.post(
            "/api/organizations/units/",
            {"khoi": self.region.id, "ten": "Đơn vị NULL", "dia_chi": None, "dien_thoai": None, "email": None, "ghi_chu": None},
            format="json",
        )
        self.assertEqual(res.status_code, status.HTTP_201_CREATED, res.data)

    def test_unauthenticated_access_denied(self):
        """Kiểm tra bắt buộc xác thực đăng nhập (401 Unauthorized)."""
        unauthenticated_client = APIClient()
        response = unauthenticated_client.get("/api/organizations/units/")
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)


class OrganizationPermissionAPITest(TestCase):
    """RBAC cho Khối và Đơn vị: 401 chưa đăng nhập, 403 thiếu quyền, 200/201/204 khi đủ quyền."""

    def setUp(self):
        self.region = KhoiKhuVuc.objects.create(ten="Khối Quyền")
        self.unit = DonVi.objects.create(khoi=self.region, stt_uu_tien=1, ten="Đơn vị Quyền")
        User = get_user_model()

        def make(username, *codenames):
            user = User.objects.create_user(username=username, password="x", don_vi=self.unit)
            if codenames:
                user.user_permissions.add(*Permission.objects.filter(codename__in=codenames))
            client = APIClient()
            client.force_authenticate(user=user)
            return client

        self.nobody = make("org_nobody")
        self.viewer = make("org_viewer", "view_donvi")
        self.admin = make("org_admin", "view_donvi", "add_donvi", "change_donvi", "delete_donvi")
        self.anon = APIClient()
        self.region_payload = {"ten": "Khối Mới"}
        self.unit_payload = {"khoi": self.region.id, "ten": "Đơn vị Mới", "is_public": True}

    def assert_denied(self, response):
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(response.data["detail"], "organizations.messages.permissionDenied")

    def test_unauthenticated_gets_401_on_every_write(self):
        for response in (
            self.anon.post("/api/organizations/regions/", self.region_payload, format="json"),
            self.anon.delete(f"/api/organizations/regions/{self.region.id}/"),
            self.anon.post("/api/organizations/units/", self.unit_payload, format="json"),
            self.anon.put(f"/api/organizations/units/{self.unit.id}/", self.unit_payload, format="json"),
            self.anon.delete(f"/api/organizations/units/{self.unit.id}/"),
        ):
            self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_user_without_permissions_is_denied_everywhere(self):
        self.assert_denied(self.nobody.get("/api/organizations/regions/"))
        self.assert_denied(self.nobody.get("/api/organizations/units/"))
        self.assert_denied(self.nobody.post("/api/organizations/regions/", self.region_payload, format="json"))
        self.assert_denied(self.nobody.post("/api/organizations/units/", self.unit_payload, format="json"))
        self.assertEqual(KhoiKhuVuc.objects.count(), 1)
        self.assertEqual(DonVi.objects.count(), 1)

    def test_view_only_user_can_read_but_not_write(self):
        self.assertEqual(self.viewer.get("/api/organizations/regions/?all=true").status_code, status.HTTP_200_OK)
        self.assertEqual(self.viewer.get("/api/organizations/units/").status_code, status.HTTP_200_OK)
        self.assertEqual(self.viewer.get(f"/api/organizations/units/{self.unit.id}/").status_code, status.HTTP_200_OK)
        self.assert_denied(self.viewer.post("/api/organizations/regions/", self.region_payload, format="json"))
        self.assert_denied(self.viewer.put(f"/api/organizations/regions/{self.region.id}/", self.region_payload, format="json"))
        self.assert_denied(self.viewer.patch(f"/api/organizations/regions/{self.region.id}/", self.region_payload, format="json"))
        self.assert_denied(self.viewer.delete(f"/api/organizations/regions/{self.region.id}/"))
        self.assert_denied(self.viewer.post("/api/organizations/units/", self.unit_payload, format="json"))
        self.assert_denied(self.viewer.put(f"/api/organizations/units/{self.unit.id}/", self.unit_payload, format="json"))
        self.assert_denied(self.viewer.patch(f"/api/organizations/units/{self.unit.id}/", {"ten": "Bị sửa"}, format="json"))
        self.assert_denied(self.viewer.delete(f"/api/organizations/units/{self.unit.id}/"))
        self.unit.refresh_from_db()
        self.assertEqual(self.unit.ten, "Đơn vị Quyền")
        self.assertEqual(KhoiKhuVuc.objects.count(), 1)

    def test_denied_before_lookup_does_not_leak_existence(self):
        self.assert_denied(self.viewer.delete("/api/organizations/units/999999/"))
        self.assert_denied(self.nobody.get("/api/organizations/regions/999999/"))

    def test_user_with_write_permissions_can_manage(self):
        created = self.admin.post("/api/organizations/regions/", self.region_payload, format="json")
        self.assertEqual(created.status_code, status.HTTP_201_CREATED)
        region_id = created.data["id"]
        self.assertEqual(
            self.admin.patch(f"/api/organizations/regions/{region_id}/", {"ten": "Khối Đổi"}, format="json").status_code,
            status.HTTP_200_OK,
        )
        unit = self.admin.post("/api/organizations/units/", {**self.unit_payload, "khoi": region_id}, format="json")
        self.assertEqual(unit.status_code, status.HTTP_201_CREATED)
        self.assertEqual(self.admin.delete(f"/api/organizations/units/{unit.data['id']}/").status_code, status.HTTP_204_NO_CONTENT)
        self.assertEqual(self.admin.delete(f"/api/organizations/regions/{region_id}/").status_code, status.HTTP_204_NO_CONTENT)

    def test_permissions_are_independent(self):
        only_add = APIClient()
        user = get_user_model().objects.create_user(username="org_only_add", password="x", don_vi=self.unit)
        user.user_permissions.add(*Permission.objects.filter(codename__in=["view_donvi", "add_donvi"]))
        only_add.force_authenticate(user=user)
        self.assertEqual(only_add.post("/api/organizations/units/", self.unit_payload, format="json").status_code, status.HTTP_201_CREATED)
        self.assert_denied(only_add.put(f"/api/organizations/units/{self.unit.id}/", self.unit_payload, format="json"))
        self.assert_denied(only_add.delete(f"/api/organizations/units/{self.unit.id}/"))

    def test_superuser_is_allowed_without_explicit_permissions(self):
        root = get_user_model().objects.create_superuser(username="org_root", password="x", email="r@example.com", don_vi=self.unit)
        client = APIClient()
        client.force_authenticate(user=root)
        self.assertEqual(client.post("/api/organizations/regions/", {"ten": "Khối Root"}, format="json").status_code, status.HTTP_201_CREATED)


class UnitAllParamAndInternalVisibilityTest(TestCase):
    """?all=true trả đủ đơn vị (không cắt ở 10); đơn vị nội bộ (is_public=False) chỉ quản trị mới thấy."""

    def setUp(self):
        self.region = KhoiKhuVuc.objects.create(ten="Khối Hiển thị")
        for index in range(1, 13):
            DonVi.objects.create(khoi=self.region, stt_uu_tien=index, ten=f"Đơn vị công khai {index}", is_public=True)
        self.internal = DonVi.objects.create(khoi=self.region, stt_uu_tien=13, ten="Đơn vị nội bộ", is_public=False)
        User = get_user_model()
        home = DonVi.objects.first()

        def make(username, *codenames, superuser=False):
            if superuser:
                user = User.objects.create_superuser(username=username, password="x", email=f"{username}@example.com", don_vi=home)
            else:
                user = User.objects.create_user(username=username, password="x", don_vi=home)
                user.user_permissions.add(*Permission.objects.filter(codename__in=codenames))
            client = APIClient()
            client.force_authenticate(user=user)
            return client

        self.viewer = make("vis_viewer", "view_donvi", "export_donvi")
        self.manager = make("vis_manager", "view_donvi", "change_donvi", "export_donvi")
        self.root = make("vis_root", superuser=True)

    def test_all_true_returns_every_visible_unit_not_only_first_page(self):
        response = self.viewer.get("/api/organizations/units/?all=true")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIsInstance(response.data, list)
        self.assertEqual(len(response.data), 12)
        paged = self.viewer.get("/api/organizations/units/")
        self.assertEqual(paged.data["count"], 12)
        self.assertEqual(len(paged.data["results"]), 10)

    def test_all_true_keeps_search_and_region_filters(self):
        found = self.viewer.get("/api/organizations/units/?all=true&search=công khai 12")
        self.assertEqual([u["ten"] for u in found.data], ["Đơn vị công khai 12"])
        other = KhoiKhuVuc.objects.create(ten="Khối khác")
        self.assertEqual(self.viewer.get(f"/api/organizations/units/?all=true&region_id={other.id}").data, [])

    def test_internal_unit_hidden_from_non_admin_everywhere(self):
        names = lambda data: [u["ten"] for u in data]
        self.assertNotIn("Đơn vị nội bộ", names(self.viewer.get("/api/organizations/units/?page_size=100").data["results"]))
        self.assertNotIn("Đơn vị nội bộ", names(self.viewer.get("/api/organizations/units/?all=true").data))
        self.assertNotIn("Đơn vị nội bộ", names(self.viewer.get("/api/organizations/units/?search=nội bộ&all=true").data))
        self.assertEqual(self.viewer.get(f"/api/organizations/units/{self.internal.id}/").status_code, status.HTTP_404_NOT_FOUND)

    def test_internal_unit_hidden_from_non_admin_export(self):
        response = self.viewer.get("/api/organizations/units/export/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
        exported = [row[1] for row in sheet.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(len(exported), 12)
        self.assertNotIn("Đơn vị nội bộ", exported)

    def test_admin_with_change_permission_and_superuser_see_internal_unit(self):
        for client in (self.manager, self.root):
            self.assertEqual(len(client.get("/api/organizations/units/?all=true").data), 13)
            self.assertEqual(client.get(f"/api/organizations/units/{self.internal.id}/").status_code, status.HTTP_200_OK)
            sheet = openpyxl.load_workbook(io.BytesIO(client.get("/api/organizations/units/export/").content)).active
            self.assertIn("Đơn vị nội bộ", [row[1] for row in sheet.iter_rows(min_row=2, values_only=True)])
