"""Xuất Excel danh bạ Phòng ban theo mẫu danh bạ: khối -> đơn vị -> mọi CBNV của đơn vị."""
import io

import openpyxl
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.db import connection
from django.db.models import Max
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from rest_framework import status
from rest_framework.test import APIClient

from accounts.models import Employee, Position
from organizations.directory_export import BLOCK_COLOR, HEADER_FILL, HEADER_ROW
from organizations.models import DonVi, KhoiKhuVuc

User = get_user_model()
URL = "/api/organizations/units/export/"


def sheet_of(response):
    return openpyxl.load_workbook(io.BytesIO(response.content)).active


def body_rows(response):
    return list(sheet_of(response).iter_rows(min_row=HEADER_ROW + 1, values_only=True))


def first_cells(response):
    return [row[0] for row in body_rows(response)]


TEST_UNITS = ("ĐƠN VỊ A1", "ĐƠN VỊ A2", "ĐƠN VỊ B1")


class DirectoryExportTest(TestCase):
    # DB test dựng bằng migration đã có 16 khối và 82 đơn vị của release_DDL.sql: các ca kiểm chỉ so phần dữ liệu
    # tự tạo, STT ưu tiên (UNIQUE) nối tiếp STT lớn nhất đang có.
    def setUp(self):
        self.block_b = KhoiKhuVuc.objects.create(ten="Khối B", thu_tu_hien_thi=2)
        self.block_a = KhoiKhuVuc.objects.create(ten="Khối A", thu_tu_hien_thi=1)
        self.priority = (DonVi.objects.aggregate(m=Max("stt_uu_tien"))["m"] or 0) + 1
        self.unit_a1 = DonVi.objects.create(
            khoi=self.block_a, stt_uu_tien=self.priority, ten="Đơn vị A1", dia_chi="1 Phố A",
            dien_thoai="0281234567", email="a1@example.com",
        )
        self.unit_a2 = DonVi.objects.create(khoi=self.block_a, stt_uu_tien=self.priority + 1, ten="Đơn vị A2")
        self.unit_b1 = DonVi.objects.create(khoi=self.block_b, stt_uu_tien=self.priority + 2, ten="Đơn vị B1")
        self.boss = Position.objects.create(display_order=101, name="Giám đốc thử")
        self.mid = Position.objects.create(display_order=102, name="Trưởng phòng thử")
        self.low = Position.objects.create(display_order=103, name="Nhân viên thử")
        # Người xuất file đứng cuối danh mục chức vụ để không thành người đứng đầu đơn vị trong các ca kiểm.
        self.lowest = Position.objects.create(display_order=999, name="Cộng tác viên thử")
        self.exporter = self.make_user("exporter", "view_donvi", "export_donvi")

    def make_user(self, username, *codenames, **extra):
        user = User.objects.create_user(
            username=username, password="x", ho_ten=username, don_vi=self.unit_b1, chuc_vu=self.lowest,
            email=f"{username}@example.com", **extra,
        )
        user.user_permissions.add(*Permission.objects.filter(codename__in=codenames))
        return user

    def person(self, username, unit, position, **extra):
        defaults = {"ho_ten": username, "chuc_danh": f"Chức danh {username}", "sdt": "0901234567"}
        defaults.update(extra)
        return User.objects.create_user(
            username=username, password="x", email=f"{username}@example.com",
            don_vi=unit, chuc_vu=position, **defaults,
        )

    def export(self, user=None, query=""):
        client = APIClient()
        client.force_authenticate(user=user or self.exporter)
        return client.get(f"{URL}{query}")

    def unit_row_index(self, rows, name):
        return next(i for i, row in enumerate(rows) if isinstance(row[0], str) and row[0].startswith(name.upper()))

    def test_title_header_and_template_colours(self):
        ws = sheet_of(self.export(query=f"?region_id={self.block_a.id}"))
        self.assertEqual(ws["A1"].value, "Bảo Long Insurance — Hệ Thống Quản Lý Văn Bản & Điều Hành")
        self.assertEqual(ws["A2"].value, "DANH BẠ ĐƠN VỊ, PHÒNG BAN")
        self.assertTrue(str(ws["A3"].value).startswith("Ngày xuất: "))
        self.assertEqual([c.value for c in ws[HEADER_ROW]], ["STT", "Họ và tên", "Chức vụ", "Chức danh", "Điện thoại", "Email"])
        self.assertEqual(ws.cell(row=HEADER_ROW, column=1).fill.start_color.rgb[-6:], HEADER_FILL)
        block_cell = ws.cell(row=HEADER_ROW + 1, column=1)
        self.assertEqual(block_cell.value, "KHỐI A")
        self.assertTrue(block_cell.font.bold)
        self.assertEqual(block_cell.font.color.rgb[-6:], BLOCK_COLOR)
        self.assertEqual(block_cell.alignment.horizontal, "center")
        unit_cell = ws.cell(row=HEADER_ROW + 2, column=1)
        self.assertTrue(unit_cell.font.bold)
        self.assertEqual(unit_cell.alignment.horizontal, "left")
        self.assertEqual(ws.freeze_panes, f"A{HEADER_ROW + 1}")
        self.assertEqual(ws.page_setup.orientation, "landscape")

    def test_blocks_follow_display_order_then_units_follow_priority(self):
        cells = first_cells(self.export())
        self.assertEqual([x for x in cells if x in ("KHỐI A", "KHỐI B")], ["KHỐI A", "KHỐI B"])
        units = [x.split("  |  ")[0] for x in cells if isinstance(x, str) and x.split("  |  ")[0] in TEST_UNITS]
        self.assertEqual(units, list(TEST_UNITS))
        self.assertLess(cells.index("KHỐI A"), cells.index(next(x for x in cells if str(x).startswith("ĐƠN VỊ A1"))))
        self.assertLess(cells.index(next(x for x in cells if str(x).startswith("ĐƠN VỊ A2"))), cells.index("KHỐI B"))

    def test_unit_row_carries_address_phone_email(self):
        label = next(x for x in first_cells(self.export()) if str(x).startswith("ĐƠN VỊ A1"))
        self.assertEqual(label, "ĐƠN VỊ A1  |  Địa chỉ: 1 Phố A  |  ĐT: 0281234567  |  Email: a1@example.com")

    def members_of(self, rows, unit_name):
        """Các dòng người ngay dưới dòng đơn vị, đến trước dòng khối hoặc đơn vị kế tiếp."""
        start = self.unit_row_index(rows, unit_name) + 1
        members = []
        for row in rows[start:]:
            if not isinstance(row[0], int):
                break
            members.append(row)
        return members

    def test_lists_every_member_by_position_order_then_lowest_id(self):
        low = self.person("nv", self.unit_a1, self.low, ho_ten="Nhân viên")
        first = self.person("tp_one", self.unit_a1, self.mid, ho_ten="Trưởng phòng Một")
        second = self.person("tp_two", self.unit_a1, self.mid, ho_ten="Trưởng phòng Hai")
        boss = self.person("gd", self.unit_a1, self.boss, ho_ten="Giám đốc")
        members = self.members_of(body_rows(self.export()), "Đơn vị A1")
        self.assertEqual([m[1] for m in members], [boss.ho_ten, first.ho_ten, second.ho_ten, low.ho_ten])
        self.assertEqual(members[1], (2, first.ho_ten, "Trưởng phòng thử", first.chuc_danh, first.sdt, first.email))

    def test_higher_position_beats_lower_id(self):
        self.person("tp_early", self.unit_a1, self.mid, ho_ten="Trưởng phòng")
        self.person("gd_late", self.unit_a1, self.boss, ho_ten="Giám đốc")
        members = self.members_of(body_rows(self.export()), "Đơn vị A1")
        self.assertEqual([m[1] for m in members], ["Giám đốc", "Trưởng phòng"])

    def test_each_person_is_listed_only_under_their_own_unit(self):
        self.person("a1", self.unit_a1, self.boss, ho_ten="Người A1")
        self.person("b1", self.unit_b1, self.boss, ho_ten="Người B1")
        rows = body_rows(self.export())
        self.assertEqual([m[1] for m in self.members_of(rows, "Đơn vị A1")], ["Người A1"])
        self.assertEqual([m[1] for m in self.members_of(rows, "Đơn vị B1")], ["Người B1", "exporter"])
        self.assertEqual(sum(1 for row in rows if row[1] == "Người A1"), 1)

    def test_serial_number_restarts_and_empty_unit_has_only_its_row(self):
        self.person("a1", self.unit_a1, self.boss)
        self.person("a1b", self.unit_a1, self.mid)
        self.person("b1", self.unit_b1, self.boss)
        rows = body_rows(self.export())
        self.assertEqual(
            [row[0] for row in self.members_of(rows, "Đơn vị A1")], [1, 2]
        )
        self.assertEqual(
            [row[0] for row in self.members_of(rows, "Đơn vị B1")], [1, 2]
        )
        self.assertEqual(self.members_of(rows, "Đơn vị A2"), [])
        self.assertEqual(rows[self.unit_row_index(rows, "Đơn vị A2") + 1][0], "KHỐI B")

    def test_resigned_staff_are_never_exported(self):
        self.person("gone", self.unit_a1, self.boss, trang_thai=Employee.Status.NGHI_VIEC, ho_ten="Đã nghỉ")
        self.person("stay", self.unit_a1, self.low, ho_ten="Còn làm")
        names = [row[1] for row in body_rows(self.export())]
        self.assertIn("Còn làm", names)
        self.assertNotIn("Đã nghỉ", names)

    def test_system_admin_account_is_not_listed(self):
        User.objects.create_superuser(
            username="sysadmin", password="x", email="sysadmin@example.com", ho_ten="Quản trị hệ thống",
            don_vi=self.unit_a1, chuc_vu=self.boss,
        )
        self.person("head", self.unit_a1, self.mid, ho_ten="Trưởng đơn vị")
        rows = body_rows(self.export())
        self.assertEqual([m[1] for m in self.members_of(rows, "Đơn vị A1")], ["Trưởng đơn vị"])
        self.assertNotIn("Quản trị hệ thống", [row[1] for row in rows])

    def test_internal_people_only_for_full_view_permission_or_superuser(self):
        self.person("hidden", self.unit_a1, self.boss, is_public=False, ho_ten="Người nội bộ")
        self.person("visible", self.unit_a1, self.low, ho_ten="Người công khai")
        plain = [row[1] for row in body_rows(self.export())]
        self.assertIn("Người công khai", plain)
        self.assertNotIn("Người nội bộ", plain)
        full = self.make_user("full", "view_donvi", "export_donvi", "view_full_cbnv")
        self.assertIn("Người nội bộ", [row[1] for row in body_rows(self.export(full))])
        root = User.objects.create_superuser(username="root", password="x", email="root@example.com", don_vi=self.unit_b1)
        self.assertIn("Người nội bộ", [row[1] for row in body_rows(self.export(root))])

    def test_filters_by_region_and_search(self):
        cells = first_cells(self.export(query=f"?region_id={self.block_b.id}"))
        self.assertEqual([x for x in cells if str(x).startswith("KHỐI")], ["KHỐI B"])
        self.assertEqual([x.split("  |  ")[0] for x in cells if str(x).startswith("ĐƠN VỊ")], ["ĐƠN VỊ B1"])
        cells = first_cells(self.export(query="?search=Đơn+vị+A2"))
        self.assertEqual([x for x in cells if str(x).startswith("KHỐI")], ["KHỐI A"])
        self.assertEqual([x.split("  |  ")[0] for x in cells if str(x).startswith("ĐƠN VỊ")], ["ĐƠN VỊ A2"])

    def test_formula_text_is_stored_as_string(self):
        self.person("evil", self.unit_a1, self.boss, ho_ten='=HYPERLINK("http://example.invalid","x")', chuc_danh="=1+1")
        ws = sheet_of(self.export())
        cells = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
        self.assertEqual(len(cells), 2)
        self.assertTrue(all(c.data_type == "s" for c in cells))

    def test_query_count_does_not_grow_with_number_of_units(self):
        self.person("head_base", self.unit_a1, self.mid)
        client = APIClient()
        client.force_authenticate(user=User.objects.get(pk=self.exporter.pk))
        client.get(URL)  # lần đầu nạp quyền của người dùng; không tính vào so sánh
        with CaptureQueriesContext(connection) as few:
            client.get(URL)
        for index in range(30):
            unit = DonVi.objects.create(khoi=self.block_a, stt_uu_tien=self.priority + 100 + index, ten=f"Đơn vị mới {index}")
            self.person(f"head{index}", unit, self.mid)
            self.person(f"staff{index}", unit, self.low)
        with CaptureQueriesContext(connection) as many:
            client.get(URL)
        self.assertEqual(len(many), len(few))

    def test_requires_export_permission(self):
        user = self.make_user("viewer_only", "view_donvi")
        response = self.export(user)
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(response.data["detail"], "organizations.messages.exportPermissionDenied")

    def test_anonymous_is_rejected(self):
        self.assertEqual(APIClient().get(URL).status_code, status.HTTP_401_UNAUTHORIZED)
