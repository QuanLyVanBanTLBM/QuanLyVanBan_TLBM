import io
from datetime import datetime, date
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from django.core.exceptions import ValidationError as DjangoValidationError
from django.http import HttpResponse
from django_filters.rest_framework import DjangoFilterBackend
from rest_framework import exceptions, filters, permissions, status, viewsets
from rest_framework.decorators import action
from rest_framework.pagination import PageNumberPagination
from rest_framework.response import Response

class StandardResourcePagination(PageNumberPagination):
    """
    Cấu trúc phân trang chuẩn REST framework: count, next, previous, results.
    Mặc định 20 bản ghi/trang, hỗ trợ tham số page_size (tối đa 100 bản ghi/trang).
    """
    page_size = 20
    page_size_query_param = 'page_size'
    max_page_size = 100


from .selectors import visible_resources
from .filters import TaiLieuChiTietFilter
from .models import DanhMucTaiLieu, LoaiTaiLieu, TaiLieuChiTiet, ChiTietTaiLieuAttachment, template_disk_path
from .serializers import (
    DanhMucTaiLieuSerializer,
    LoaiTaiLieuSerializer,
    TaiLieuChiTietSerializer,
    ChiTietTaiLieuAttachmentSerializer,
)


class DanhMucTaiLieuViewSet(viewsets.ModelViewSet):
    """
    ViewSet cho Danh mục Tài liệu (Cấp 1):
    - CRUD danh mục
    - Tự động gán người lập nếu đã đăng nhập
    - Tìm kiếm theo tên tài liệu, số ký hiệu
    - Xuất Excel danh mục Cấp 1
    """
    queryset = DanhMucTaiLieu.objects.select_related('nguoi_lap').all()
    serializer_class = DanhMucTaiLieuSerializer
    permission_classes = [permissions.IsAuthenticatedOrReadOnly]
    pagination_class = StandardResourcePagination
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['ten_tai_lieu', 'ky_hieu', 'ghi_chu']
    ordering_fields = ['stt_uu_tien', 'id', 'created_at', 'ten_tai_lieu']
    ordering = ['stt_uu_tien', 'id']

    def perform_create(self, serializer):
        kwargs = {}
        if self.request.user and self.request.user.is_authenticated:
            if 'nguoi_lap' not in serializer.validated_data:
                kwargs['nguoi_lap'] = self.request.user
        serializer.save(**kwargs)

    @action(detail=False, methods=['get'], url_path='export-excel')
    def export_excel(self, request):
        qs = self.filter_queryset(self.get_queryset())
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DanhMucCap1"

        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin", color="CCCCCC"),
            right=Side(style="thin", color="CCCCCC"),
            top=Side(style="thin", color="CCCCCC"),
            bottom=Side(style="thin", color="CCCCCC"),
        )

        headers = ["STT", "Số ký hiệu", "Tên tài liệu / Danh mục", "Người lập", "Ngày lập", "Ghi chú"]
        ws.append(headers)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align

        for idx, item in enumerate(qs, start=1):
            nguoi_lap_str = (
                item.nguoi_lap.get_full_name() or item.nguoi_lap.username
                if item.nguoi_lap else ""
            )
            ngay_lap_str = item.ngay_lap.strftime("%d/%m/%Y") if item.ngay_lap else ""
            row_data = [
                item.stt_uu_tien or idx,
                item.ky_hieu or "",
                item.ten_tai_lieu,
                nguoi_lap_str,
                ngay_lap_str,
                item.ghi_chu or "",
            ]
            ws.append(row_data)

        # Style data rows and auto width
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
            for cell in row:
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        date_str = datetime.now().strftime("%Y%m%d_%H%M%S")
        response["Content-Disposition"] = f'attachment; filename="Danh_muc_cap1_{date_str}.xlsx"'
        return response


class LoaiTaiLieuViewSet(viewsets.ModelViewSet):
    """
    ViewSet cho Loại tài liệu (Cấp 2):
    - CRUD loại tài liệu
    - Lọc theo Cấp 1 (`dm_tai_lieu`)
    - Tìm kiếm theo tên loại
    """
    queryset = LoaiTaiLieu.objects.select_related('dm_tai_lieu').all()
    serializer_class = LoaiTaiLieuSerializer
    permission_classes = [permissions.IsAuthenticatedOrReadOnly]
    pagination_class = StandardResourcePagination
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = ['dm_tai_lieu', 'is_public']
    search_fields = ['ten_loai', 'ghi_chu']
    ordering_fields = ['id', 'ten_loai', 'created_at']
    ordering = ['id']


class TaiLieuChiTietViewSet(viewsets.ModelViewSet):
    """
    ViewSet cho Chi tiết tài liệu chung:
    - CRUD chi tiết tài liệu (kèm ManyToMany loai_tai_lieu và FileField upload)
    - Quản lý trạng thái hiệu lực: CON_HIEU_LUC, HET_HIEU_LUC, HET_VA_THAY_THE
    - Upload file đính kèm với validator (< 20MB, định dạng hợp lệ)
    - Phân quyền hiển thị: TAT_CA hoặc NOI_BO_PHONG_BAN
    - Bộ lọc đa tiêu chí: dm_tai_lieu, loai_tai_lieu, nam, hieu_luc, search
    - Xuất Excel danh sách tài liệu theo đúng bộ lọc tìm kiếm
    """
    serializer_class = TaiLieuChiTietSerializer
    permission_classes = [permissions.IsAuthenticatedOrReadOnly]
    pagination_class = StandardResourcePagination
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_class = TaiLieuChiTietFilter
    search_fields = ['tieu_de', 'ky_hieu', 'noi_ban_hanh', 'ghi_chu']
    ordering_fields = ['ngay_dang', 'id', 'tieu_de', 'ky_hieu']
    ordering = ['-ngay_dang', '-id']

    def get_queryset(self):
        qs = TaiLieuChiTiet.objects.select_related(
            'dm_tai_lieu', 'nguoi_tao'
        ).prefetch_related('loai_tai_lieu', 'attachments').all()

        return visible_resources(qs, self.request.user)

    def perform_create(self, serializer):
        kwargs = {}
        if self.request.user and self.request.user.is_authenticated:
            kwargs['nguoi_tao'] = self.request.user
            if 'don_vi_id' not in serializer.validated_data and hasattr(self.request.user, 'don_vi_id'):
                kwargs['don_vi_id'] = getattr(self.request.user, 'don_vi_id')
        serializer.save(**kwargs)

    @action(detail=True, methods=['post'], url_path='toggle-validity')
    def toggle_validity(self, request, pk=None):
        """
        API chuyển đổi nhanh trạng thái hiệu lực
        """
        obj = self.get_object()
        target_status = request.data.get('hieu_luc') or request.data.get('status')
        link_thay_the = request.data.get('link_thay_the')
        ngay_het_hieu_luc = request.data.get('ngay_het_hieu_luc')

        if target_status:
            valid_statuses = [choice[0] for choice in TaiLieuChiTiet.HIEU_LUC_CHOICES]
            if target_status not in valid_statuses:
                return Response(
                    {'error': f"Trạng thái '{target_status}' không hợp lệ. Cho phép: {valid_statuses}"},
                    status=status.HTTP_400_BAD_REQUEST
                )
            obj.hieu_luc = target_status
        else:
            # Toggle giữa CON_HIEU_LUC và HET_HIEU_LUC
            if obj.hieu_luc == 'CON_HIEU_LUC':
                obj.hieu_luc = 'HET_HIEU_LUC'
            else:
                obj.hieu_luc = 'CON_HIEU_LUC'

        if obj.hieu_luc == 'CON_HIEU_LUC':
            obj.link_thay_the = ''
            if ngay_het_hieu_luc is not None:
                obj.ngay_het_hieu_luc = ngay_het_hieu_luc if ngay_het_hieu_luc != '' else None
            else:
                obj.ngay_het_hieu_luc = None
        else:
            if link_thay_the is not None:
                obj.link_thay_the = link_thay_the
            if ngay_het_hieu_luc is not None:
                obj.ngay_het_hieu_luc = ngay_het_hieu_luc if ngay_het_hieu_luc != '' else None

        try:
            obj.full_clean()
            obj.save()
        except DjangoValidationError as e:
            return Response(
                e.message_dict if hasattr(e, 'message_dict') else {'error': str(e)},
                status=status.HTTP_400_BAD_REQUEST
            )

        serializer = self.get_serializer(obj)
        return Response(serializer.data)

    @action(detail=False, methods=['get'], url_path='export-excel')
    def export_excel(self, request):
        """
        Xuất danh sách tài liệu chi tiết sang file Excel (.xlsx) theo đúng bộ lọc tìm kiếm hiện tại
        """
        filtered_qs = self.filter_queryset(self.get_queryset())
        lang = request.query_params.get('lang', 'vi')
        is_en = (lang == 'en')

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DocumentList" if is_en else "DanhSachTaiLieu"

        # Cấu hình phong cách bảng tính
        title_font = Font(name="Arial", size=14, bold=True, color="1E3A8A")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        dimmed_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        dimmed_font = Font(name="Arial", size=10, color="64748B")
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin", color="CCCCCC"),
            right=Side(style="thin", color="CCCCCC"),
            top=Side(style="thin", color="CCCCCC"),
            bottom=Side(style="thin", color="CCCCCC"),
        )

        # Tiêu đề bảng
        ws.merge_cells("A1:L1")
        title_cell = ws["A1"]
        title_cell.value = (
            "COMMON DOCUMENTS AND TEMPLATES DIRECTORY"
            if is_en
            else "DANH SÁCH TÀI LIỆU VÀ BIỂU MẪU CHUNG"
        )
        title_cell.font = title_font
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 35

        # Dòng thời gian xuất
        ws.merge_cells("A2:L2")
        sub_cell = ws["A2"]
        if is_en:
            sub_cell.value = f"Export date: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Total documents: {filtered_qs.count()}"
        else:
            sub_cell.value = f"Ngày xuất báo cáo: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Tổng số tài liệu: {filtered_qs.count()}"
        sub_cell.font = Font(name="Arial", size=10, italic=True, color="475569")
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20

        # Header
        if is_en:
            headers = [
                "No.",
                "Document Code",
                "Document Title",
                "Category (Level 1)",
                "Document Type (Level 2)",
                "Issuing Authority",
                "Published Date",
                "Status",
                "Expiration Date",
                "Replacement Document",
                "Visibility Scope",
                "Created By",
            ]
        else:
            headers = [
                "STT",
                "Số ký hiệu",
                "Tiêu đề tài liệu",
                "Danh mục Cấp 1",
                "Loại tài liệu Cấp 2",
                "Nơi ban hành",
                "Ngày đăng",
                "Hiệu lực",
                "Ngày hết hiệu lực",
                "Văn bản thay thế",
                "Phạm vi hiển thị",
                "Người tạo",
            ]
        ws.append([])  # Dòng trống row 3
        ws.append(headers)  # Row 4
        ws.row_dimensions[4].height = 28

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=4, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        # Điền dữ liệu
        for idx, doc in enumerate(filtered_qs, start=1):
            loai_names = ", ".join([lt.ten_loai for lt in doc.loai_tai_lieu.all()])
            ngay_dang_str = doc.ngay_dang.strftime("%d/%m/%Y") if doc.ngay_dang else ""
            ngay_het_str = doc.ngay_het_hieu_luc.strftime("%d/%m/%Y") if doc.ngay_het_hieu_luc else ""
            nguoi_tao_str = (
                doc.nguoi_tao.get_full_name() or doc.nguoi_tao.username
                if doc.nguoi_tao else ""
            )

            is_expired = doc.is_expired
            effective_status = doc.effective_hieu_luc

            if is_en:
                hieu_luc_map = {
                    'CON_HIEU_LUC': 'Active',
                    'CON_HIEU_LUC_1_PHAN': 'Partially Active',
                    'HET_HIEU_LUC': 'Expired',
                    'HET_VA_THAY_THE': 'Expired & Replaced',
                }
                hieu_luc_display = hieu_luc_map.get(effective_status, effective_status)
                hien_thi_map = {
                    'TAT_CA': 'Public / All',
                    'PHONG_BAN': 'Internal Department',
                }
                hien_thi_display = hien_thi_map.get(doc.hien_thi, doc.get_hien_thi_display())
            else:
                hieu_luc_vn_map = dict(TaiLieuChiTiet.HIEU_LUC_CHOICES)
                hieu_luc_display = hieu_luc_vn_map.get(effective_status, effective_status)
                hien_thi_display = doc.get_hien_thi_display()

            row_data = [
                idx,
                doc.ky_hieu or "",
                doc.tieu_de,
                doc.dm_tai_lieu.ten_tai_lieu if doc.dm_tai_lieu else "",
                loai_names,
                doc.noi_ban_hanh or "",
                ngay_dang_str,
                hieu_luc_display,
                ngay_het_str,
                doc.link_thay_the or "",
                hien_thi_display,
                nguoi_tao_str,
            ]
            ws.append(row_data)
            current_row = ws.max_row
            ws.row_dimensions[current_row].height = 24

            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=current_row, column=col_idx)
                c.border = thin_border
                if is_expired:
                    c.fill = dimmed_fill
                    c.font = dimmed_font

                # Căn lề
                if col_idx in [1, 2, 7, 8, 9, 11]:
                    c.alignment = center_align
                else:
                    c.alignment = left_align

        # Tự động canh chỉnh độ rộng cột
        for col in ws.columns:
            if col[0].row > 4:
                continue
            max_len = 0
            for cell in col:
                if cell.row < 4:
                    continue
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        date_str = datetime.now().strftime("%Y%m%d_%H%M%S")
        prefix = "Document_List" if is_en else "Danh_sach_tai_lieu"
        response["Content-Disposition"] = f'attachment; filename="{prefix}_{date_str}.xlsx"'
        return response


class ChiTietTaiLieuAttachmentViewSet(viewsets.ModelViewSet):
    """
    ViewSet cho Tệp đính kèm tài liệu (1-N Hướng A):
    - Quản lý danh sách, upload, xóa mềm tệp đính kèm theo DDL Base
    - Chống lỗ hổng IDOR rò rỉ tệp nội bộ: lọc tệp theo visible_resources của tài liệu cha
    - Kiểm soát chặt quyền xóa mềm và upload tệp
    """
    serializer_class = ChiTietTaiLieuAttachmentSerializer
    permission_classes = [permissions.IsAuthenticatedOrReadOnly]
    pagination_class = StandardResourcePagination
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = ['document']
    search_fields = ['name']
    ordering_fields = ['id', 'created_at', 'name', 'file_size']
    ordering = ['id']

    def get_queryset(self):
        user = self.request.user
        # Bảo vệ IDOR: Chỉ cho phép truy vấn tệp thuộc về các tài liệu mà người dùng có quyền xem
        visible_docs = visible_resources(TaiLieuChiTiet.objects.all(), user)
        return ChiTietTaiLieuAttachment.objects.filter(
            is_deleted=False,
            document__in=visible_docs
        ).select_related('document', 'nguoi_tao')

    def perform_create(self, serializer):
        user = self.request.user
        if not user or not user.is_authenticated:
            raise exceptions.PermissionDenied("Yêu cầu đăng nhập để tải lên tệp đính kèm.")
        doc = serializer.validated_data.get('document')
        can_add = (
            user.is_superuser or
            user.is_staff or
            (doc and doc.nguoi_tao == user) or
            (doc and hasattr(user, 'don_vi_id') and user.don_vi_id and doc.don_vi_id == user.don_vi_id)
        )
        if not can_add:
            raise exceptions.PermissionDenied("Không có quyền tải lên tệp đính kèm cho tài liệu này.")
        serializer.save(nguoi_tao=user)

    def perform_destroy(self, instance):
        user = self.request.user
        if not user or not user.is_authenticated:
            raise exceptions.PermissionDenied("Yêu cầu đăng nhập để xóa tệp đính kèm.")
        can_delete = (
            user.is_superuser or
            user.is_staff or
            instance.nguoi_tao == user or
            instance.document.nguoi_tao == user or
            (hasattr(user, 'don_vi_id') and user.don_vi_id and instance.document.don_vi_id == user.don_vi_id)
        )
        if not can_delete:
            raise exceptions.PermissionDenied("Không có quyền xóa tệp đính kèm này.")
        instance.is_deleted = True
        instance.save()

    @action(detail=True, methods=['get'], url_path='view')
    def view_file(self, request, pk=None):
        """
        Mở xem trực tiếp tệp trong tab trình duyệt (inline display, không bắt buộc tải về)
        """
        import os
        import mimetypes
        from django.http import FileResponse, Http404

        att = self.get_object()
        disk_path = template_disk_path(att.path)
        if not disk_path or not os.path.exists(disk_path):
            raise Http404("Tệp không tồn tại trên máy chủ")

        content_type, _ = mimetypes.guess_type(disk_path)
        if not content_type:
            ext = att.name.split('.')[-1].lower() if '.' in att.name else ''
            if ext == 'pdf':
                content_type = 'application/pdf'
            elif ext in ['xlsx', 'xls']:
                content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' if ext == 'xlsx' else 'application/vnd.ms-excel'
            elif ext in ['png', 'jpg', 'jpeg']:
                content_type = f'image/{ext}'
            else:
                content_type = 'application/octet-stream'

        response = FileResponse(open(disk_path, 'rb'), content_type=content_type, as_attachment=False)
        response['Content-Disposition'] = f'inline; filename="{att.name}"'
        return response

    @action(detail=True, methods=['get'], url_path='download')
    def download_file(self, request, pk=None):
        """
        Tải tệp đính kèm về máy (as_attachment=True)
        """
        import os
        import mimetypes
        from django.http import FileResponse, Http404

        att = self.get_object()
        disk_path = template_disk_path(att.path)
        if not disk_path or not os.path.exists(disk_path):
            raise Http404("Tệp không tồn tại trên máy chủ")

        content_type, _ = mimetypes.guess_type(disk_path)
        if not content_type:
            content_type = 'application/octet-stream'

        response = FileResponse(open(disk_path, 'rb'), content_type=content_type, as_attachment=True, filename=att.name)
        return response

    @action(detail=True, methods=['get'], url_path='preview')
    def preview(self, request, pk=None):
        """
        API xem trước chi tiết tệp đính kèm:
        - Excel (.xlsx, .xls): đọc dữ liệu các sheet trả về mảng bảng tính JSON
        - PDF: trả về thông tin và URL trực tiếp
        - Image: trả về thông tin và URL trực tiếp
        - Khác: trả về thông tin tệp
        """
        att = self.get_object()
        file_name = att.name
        ext = file_name.split('.')[-1].lower() if '.' in file_name else ''
        file_url = att.path

        if ext in ['xlsx', 'xls']:
            sheets_data = []
            try:
                import os
                disk_path = template_disk_path(file_url)
                if disk_path and os.path.exists(disk_path):
                    wb = openpyxl.load_workbook(disk_path, data_only=True)
                    for s_name in wb.sheetnames[:5]:
                        ws = wb[s_name]
                        rows_data = []
                        for row in ws.iter_rows(max_row=100, max_col=26, values_only=True):
                            row_vals = [
                                cell.strftime('%d/%m/%Y') if isinstance(cell, (datetime, date))
                                else str(cell) if cell is not None
                                else ""
                                for cell in row
                            ]
                            if any(v != "" for v in row_vals):
                                rows_data.append(row_vals)
                        sheets_data.append({
                            'name': s_name,
                            'rows': rows_data
                        })
                    wb.close()
            except Exception:
                pass
            return Response({
                'id': att.id,
                'name': att.name,
                'file_type': 'EXCEL',
                'file_size': att.file_size,
                'file_url': file_url,
                'sheets': sheets_data
            })

        if ext == 'pdf':
            return Response({
                'id': att.id,
                'name': att.name,
                'file_type': 'PDF',
                'file_size': att.file_size,
                'file_url': file_url,
            })

        if ext in ['png', 'jpg', 'jpeg', 'webp', 'gif']:
            return Response({
                'id': att.id,
                'name': att.name,
                'file_type': 'IMAGE',
                'file_size': att.file_size,
                'file_url': file_url,
            })

        return Response({
            'id': att.id,
            'name': att.name,
            'file_type': ext.upper() or 'DOCUMENT',
            'file_size': att.file_size,
            'file_url': file_url,
        })


